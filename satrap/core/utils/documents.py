"""
文档解析: xlsx / docx / pdf / 纯文本 提取为纯文本

按扩展名分发到对应解析器, 统一返回纯文本; 解析失败抛出带明确信息的异常
管理接口与插件共用解析器, Office 和 PDF 解析依赖按需加载
"""
from __future__ import annotations

from collections.abc import Iterable, Sequence
from importlib.util import find_spec
import logging
from pathlib import Path
from zipfile import BadZipFile, ZipFile
from typing import Protocol, cast

_TEXT_EXTS = {".txt", ".md", ".py", ".json", ".csv", ".log", ".yaml", ".yml", ".toml", ".xml", ".html", ".js", ".ts"}
# 纯文本可直接读取的扩展名

SUPPORTED_EXTENSIONS = tuple(sorted(_TEXT_EXTS | {".pdf", ".docx", ".xlsx"}))
_PARSER_MODULES = {".pdf": "pdfplumber", ".docx": "docx", ".xlsx": "openpyxl"}

DEFAULT_MAX_FILE_SIZE = 32 * 1024 * 1024
DEFAULT_MAX_EXPANDED_SIZE = 128 * 1024 * 1024
DEFAULT_MAX_COMPRESSION_RATIO = 200.0
DEFAULT_MAX_PDF_PAGES = 500
DEFAULT_MAX_XLSX_ROWS = 100_000
DEFAULT_MAX_XLSX_CELLS = 1_000_000
DEFAULT_MAX_OUTPUT_LENGTH = 1_000_000


class _Worksheet(Protocol):
    """声明工作表解析所需的最小接口"""

    title: str

    def iter_rows(self, *, values_only: bool) -> Iterable[tuple[object, ...]]:
        """
        按行迭代单元格值

        参数:
        - values_only: 是否仅返回单元格值

        返回:
        - 工作表行迭代器
        """
        ...


class _Workbook(Protocol):
    """声明工作簿解析所需的最小接口"""

    worksheets: Sequence[_Worksheet]

    def close(self) -> None:
        """关闭工作簿资源"""
        ...


class _OpenpyxlModule(Protocol):
    """声明 openpyxl 加载工作簿的最小接口"""

    def load_workbook(
        self,
        filename: str,
        *,
        read_only: bool,
        data_only: bool,
    ) -> _Workbook:
        """
        加载只读工作簿

        参数:
        - filename: 工作簿路径
        - read_only: 是否按只读模式加载
        - data_only: 是否读取公式计算值

        返回:
        - 工作簿读取接口
        """
        ...


class _TextBudget:
    """解析阶段的文本输出预算"""

    def __init__(self, maximum: int) -> None:
        """
        初始化文本输出预算

        参数:
        - maximum: 最多保留的字符数
        """
        self.maximum = max(1, int(maximum))
        self.length = 0
        self.parts: list[str] = []

    def append(self, text: str) -> bool:
        """
        追加预算内文本

        参数:
        - text: 待追加文本

        返回:
        - 文本完整写入预算时返回 True; 被截断时返回 False
        """
        remaining = self.maximum - self.length
        if remaining <= 0:
            return False
        chunk = text[:remaining]
        self.parts.append(chunk)
        self.length += len(chunk)
        return len(text) <= remaining

    def render(self) -> str:
        """
        返回已收集文本

        返回:
        - 按追加顺序拼接的预算内文本
        """
        return "".join(self.parts)


def _validate_zip_limits(path: Path, max_expanded_size: int, max_compression_ratio: float) -> None:
    """
    校验 OOXML 压缩包的展开大小与压缩比

    参数:
    - path: OOXML 文档路径
    - max_expanded_size: 允许的最大展开字节数
    - max_compression_ratio: 允许的最大压缩比
    """
    try:
        with ZipFile(path) as archive:
            infos = archive.infolist()
            expanded = sum(info.file_size for info in infos)
            compressed = sum(info.compress_size for info in infos)
    except BadZipFile as error:
        raise ValueError(f"文档压缩包损坏: {error}") from error
    if expanded > max_expanded_size:
        raise ValueError(f"文档展开后过大: {expanded} 字节, 上限 {max_expanded_size} 字节")
    if expanded / max(1, compressed) > max_compression_ratio:
        raise ValueError(f"文档压缩比过高, 上限 {max_compression_ratio:g}")


def _read_xlsx(path: Path, max_length: int, max_rows: int, max_cells: int) -> str:
    """
    openpyxl 读所有 sheet, 每行拼 TSV 文本

    参数:
    - path: 路径
    - max_length: 最多返回的文本字符数
    - max_rows: 最多解析的总行数
    - max_cells: 最多解析的总单元格数

    返回:
    - str: openpyxl 读所有 sheet, 每行拼 TSV 文本
    """
    import openpyxl   # 仅在读取 xlsx 时加载可选依赖
    workbook = cast(_OpenpyxlModule, openpyxl).load_workbook(
        str(path), read_only=True, data_only=True,
    )
    budget = _TextBudget(max_length)
    row_count = 0
    cell_count = 0
    try:
        for sheet in workbook.worksheets:
            heading_written = False
            for row in sheet.iter_rows(values_only=True):
                row_count += 1
                cell_count += len(row)
                if row_count > max_rows or cell_count > max_cells:
                    raise ValueError(
                        f"工作簿内容超出解析上限: 行 {max_rows}, 单元格 {max_cells}"
                    )
                cells = ["" if cell is None else str(cell) for cell in row]
                if any(cell.strip() for cell in cells):
                    if not heading_written:
                        if not budget.append(f"# Sheet: {sheet.title}\n"):
                            break
                        heading_written = True
                    if not budget.append("\t".join(cells).rstrip() + "\n"):
                        break
            if budget.length >= budget.maximum:
                break
    finally:
        workbook.close()
    return budget.render().rstrip()


def _read_docx(path: Path, max_length: int) -> str:
    """
    python-docx 读段落 + 表格, 拼纯文本

    参数:
    - path: 路径
    - max_length: 最多返回的文本字符数

    返回:
    - str: python-docx 读段落 + 表格, 拼纯文本
    """
    import docx   # 仅在读取 docx 时加载可选依赖

    doc = docx.Document(str(path))
    budget = _TextBudget(max_length)
    for para in doc.paragraphs:
        text = para.text.strip()
        if text and not budget.append(text + "\n"):
            return budget.render().rstrip()
    for ti, table in enumerate(doc.tables, 1):
        heading_written = False
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if not any(cells):
                continue
            if not heading_written:
                if not budget.append(f"# Table {ti}\n"):
                    return budget.render().rstrip()
                heading_written = True
            if not budget.append("\t".join(cells).rstrip() + "\n"):
                return budget.render().rstrip()
    return budget.render().rstrip()


class _FontBBoxWarningFilter(logging.Filter):
    """过滤 pdfminer 对缺 FontBBox 字体的警告: 纯文本提取不依赖 bbox, 属良性噪音"""

    _MESSAGE = "Could not get FontBBox from font descriptor"

    def filter(self, record: logging.LogRecord) -> bool:
        """
        执行 `filter` 操作

        参数:
        - record: 记录

        返回:
        - bool: 执行 `filter` 操作
        """
        try:
            msg = record.getMessage()
        except Exception:
            return True
        return self._MESSAGE not in msg


def _mute_pdfminer_fontbbox_warning() -> None:
    """给 pdfminer.pdffont 挂一次性过滤器, 仅吞掉上述良性警告 (幂等)"""
    logger = logging.getLogger("pdfminer.pdffont")
    if any(isinstance(f, _FontBBoxWarningFilter) for f in logger.filters):
        return
    logger.addFilter(_FontBBoxWarningFilter())


def _read_pdf(path: Path, max_length: int, max_pages: int) -> str:
    """
    pdfplumber 逐页 extract_text

    参数:
    - path: 路径
    - max_length: 最多返回的文本字符数
    - max_pages: 最多解析的 PDF 页数

    返回:
    - str: pdfplumber 逐页 extract_text
    """
    import pdfplumber   # 仅在读取 pdf 时加载可选依赖

    _mute_pdfminer_fontbbox_warning()

    budget = _TextBudget(max_length)
    with pdfplumber.open(str(path)) as pdf:
        if len(pdf.pages) > max_pages:
            raise ValueError(f"PDF 页数过多: {len(pdf.pages)}, 上限 {max_pages}")
        for i, page in enumerate(pdf.pages, 1):
            text = page.extract_text() or ""
            if not text.strip():
                continue
            if not budget.append(f"# Page {i}\n{text.strip()}\n\n"):
                break
    return budget.render().rstrip()


def extract_text(
    path: str | Path,
    *,
    max_length: int = DEFAULT_MAX_OUTPUT_LENGTH,
    max_file_size: int = DEFAULT_MAX_FILE_SIZE,
    max_expanded_size: int = DEFAULT_MAX_EXPANDED_SIZE,
    max_compression_ratio: float = DEFAULT_MAX_COMPRESSION_RATIO,
    max_pdf_pages: int = DEFAULT_MAX_PDF_PAGES,
    max_xlsx_rows: int = DEFAULT_MAX_XLSX_ROWS,
    max_xlsx_cells: int = DEFAULT_MAX_XLSX_CELLS,
) -> str:
    """
    按扩展名提取文档纯文本

    参数:
    - path: 路径
    - max_length: 最多返回的文本字符数
    - max_file_size: 允许的最大源文件字节数
    - max_expanded_size: OOXML 文档允许的最大展开字节数
    - max_compression_ratio: OOXML 文档允许的最大压缩比
    - max_pdf_pages: 最多解析的 PDF 页数
    - max_xlsx_rows: 最多解析的工作簿总行数
    - max_xlsx_cells: 最多解析的工作簿总单元格数

    支持: .xlsx / .docx / .pdf / 常见纯文本 (.txt/.md/.py 等)
    解析失败抛出 ValueError (带明确原因)

    返回:
    - str: 按扩展名提取文档纯文本
    """
    p = Path(path)
    if not p.is_file():
        raise ValueError(f"文件不存在: {p}")
    file_size = p.stat().st_size
    if file_size > max_file_size:
        raise ValueError(f"文件过大: {file_size} 字节, 上限 {max_file_size} 字节")
    ext = p.suffix.lower()
    try:
        if ext == ".xlsx":
            _validate_zip_limits(p, max_expanded_size, max_compression_ratio)
            return _read_xlsx(p, max_length, max_xlsx_rows, max_xlsx_cells)
        if ext == ".docx":
            _validate_zip_limits(p, max_expanded_size, max_compression_ratio)
            return _read_docx(p, max_length)
        if ext == ".pdf":
            return _read_pdf(p, max_length, max_pdf_pages)
        if ext in _TEXT_EXTS:
            with p.open("r", encoding="utf-8-sig", errors="strict") as file:
                return file.read(max(1, int(max_length)))
    except UnicodeDecodeError as error:
        raise ValueError("文本不是有效的 UTF-8 编码, 请转换编码后重试") from error
    except ValueError:
        raise
    except Exception as e:
        raise ValueError(f"解析 {ext} 失败: {e}") from e
    raise ValueError(f"不支持的文档类型: {ext or '(无扩展名)'} (支持 xlsx/docx/pdf/txt/md 等)")


def document_capabilities() -> dict[str, object]:
    """
    返回前后端共享的上传约束及可选解析器状态

    返回:
    - 扩展名, 文件及文本大小限制, 缺失的解析依赖
    """
    missing = {ext: module for ext, module in _PARSER_MODULES.items() if find_spec(module) is None}
    return {
        "extensions": list(SUPPORTED_EXTENSIONS),
        "max_file_bytes": DEFAULT_MAX_FILE_SIZE,
        "max_text_chars": DEFAULT_MAX_OUTPUT_LENGTH,
        "missing_parsers": missing,
        "text_encoding": "UTF-8",
    }


def extract_document(path: str | Path, *, max_length: int = DEFAULT_MAX_OUTPUT_LENGTH) -> str:
    """
    完整提取入库文档, 拒绝乱码, 空文档和超限内容

    参数:
    - path: 待解析的本地文件
    - max_length: 允许的最大文本字符数, 默认 1000000

    返回:
    - 非空文档文本, 无法完整提取时抛出 ValueError
    """
    if max_length < 1:
        raise ValueError("文本字符上限必须为正整数")
    suffix = Path(path).suffix.lower()
    module = _PARSER_MODULES.get(suffix)
    if module and find_spec(module) is None:
        raise ValueError(f"无法解析 {suffix}, 缺少解析依赖 {module}")
    text = extract_text(path, max_length=max_length + 1)
    if len(text) > max_length:
        raise ValueError(f"文档文本超过 {max_length} 字符, 请拆分后导入")
    if not text.strip():
        message = "文档没有可提取文字"
        if suffix == ".pdf":
            message += ", 扫描 PDF 暂不支持 OCR"
        raise ValueError(message)
    if "\x00" in text:
        raise ValueError("文档包含空字节, 请使用 UTF-8 文本")
    return text
