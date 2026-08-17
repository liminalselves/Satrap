"""文档解析: xlsx / docx / pdf / 纯文本 提取为纯文本

按扩展名分发到对应解析器, 统一返回纯文本; 解析失败抛出带明确信息的异常
依赖: openpyxl (xlsx) / python-docx (docx) / pdfplumber (pdf), 均入 requirements
"""
from __future__ import annotations

from pathlib import Path

# 纯文本可直接读取的扩展名
_TEXT_EXTS = {".txt", ".md", ".py", ".json", ".csv", ".log", ".yaml", ".yml", ".toml", ".xml", ".html", ".js", ".ts"}


def _read_xlsx(path: Path) -> str:
    """openpyxl 读所有 sheet, 每行拼 TSV 文本"""
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    parts: list[str] = []
    for sheet in wb.worksheets:
        parts.append(f"# Sheet: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            cells = ["" if c is None else str(c) for c in row]
            # 跳过全空行
            if any(c.strip() for c in cells):
                parts.append("\t".join(cells).rstrip())
    wb.close()
    return "\n".join(parts)


def _read_docx(path: Path) -> str:
    """python-docx 读段落 + 表格, 拼纯文本"""
    import docx

    doc = docx.Document(str(path))
    parts: list[str] = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            parts.append(text)
    for ti, table in enumerate(doc.tables, 1):
        parts.append(f"# Table {ti}")
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            parts.append("\t".join(cells).rstrip())
    return "\n".join(parts)


def _read_pdf(path: Path) -> str:
    """pdfplumber 逐页 extract_text"""
    import pdfplumber

    parts: list[str] = []
    with pdfplumber.open(str(path)) as pdf:
        for i, page in enumerate(pdf.pages, 1):
            text = page.extract_text() or ""
            parts.append(f"# Page {i}\n{text.strip()}")
    return "\n\n".join(parts)


def extract_text(path: str | Path) -> str:
    """按扩展名提取文档纯文本

    支持: .xlsx / .docx / .pdf / 常见纯文本 (.txt/.md/.py 等)
    解析失败抛出 ValueError (带明确原因)
    """
    p = Path(path)
    if not p.is_file():
        raise ValueError(f"文件不存在: {p}")
    ext = p.suffix.lower()
    try:
        if ext == ".xlsx":
            return _read_xlsx(p)
        if ext == ".docx":
            return _read_docx(p)
        if ext == ".pdf":
            return _read_pdf(p)
        if ext in _TEXT_EXTS:
            return p.read_text(encoding="utf-8", errors="replace")
    except ValueError:
        raise
    except Exception as e:
        raise ValueError(f"解析 {ext} 失败: {e}") from e
    raise ValueError(f"不支持的文档类型: {ext or '(无扩展名)'} (支持 xlsx/docx/pdf/txt/md 等)")
