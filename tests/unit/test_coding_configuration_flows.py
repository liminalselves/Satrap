"""业务链路场景与集中边界契约"""
from __future__ import annotations

import inspect
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock
from typing import Any
import pytest
from satrap.expend.plugins.satrap_coding import tools as coding
from satrap.expend.plugins.satrap_coding.commands import build_commands
from satrap.expend.plugins.satrap_coding.handlers import build_handlers
from satrap.expend.plugins.satrap_coding.state import get_plugin_state, reset_plugin_state
from satrap.expend.plugins.satrap_coding.tools import async_interaction, sync_interaction
from satrap.core.APICall.LLMCall import AsyncLLM, LLM
from satrap.edictum import AsyncSimpleSession, HandlerConfig, HandlerContext, SimpleSession
from typing import Iterator, cast
from satrap.expend.plugins.satrap_coding import tools as tools_mod
from satrap.core.type import LLMCallResponse, LLMCallStreamEvent
from zipfile import ZIP_DEFLATED, ZipFile
from typing import Callable, Protocol
from satrap.expend.plugins.base_take.core.docread import extract_text

_coding_configuration_PLUGIN_DIR = Path(__file__).resolve().parents[2] / "satrap/expend/plugins/satrap_coding"

async def _execute(tool: Any, *args: Any, **kwargs: Any) -> str:
    """执行同步或异步测试工具"""
    result = tool.execute(*args, **kwargs)
    return await result if inspect.isawaitable(result) else result

@pytest.mark.asyncio
@pytest.mark.parametrize("asynchronous", [False, True], ids=["sync", "async"])
async def test_factory_configuration_isolated_without_host_paths(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, asynchronous: bool,
) -> None:
    """无宿主路径时, 交错构造和修改原配置均不改变工具路径与默认超时"""
    def embedded_session(name: str) -> Any:
        session = object.__new__(AsyncSimpleSession if asynchronous else SimpleSession)
        session.session_id = name
        setattr(session, "_wf", SimpleNamespace(llm=None, tools_manager=None))
        setattr(session, "session_ctx", SimpleNamespace(db_path=str(tmp_path / "chat.db")))
        return session

    sessions = [embedded_session(name) for name in ("a", "b", "empty", "none")]
    configs: list[dict[str, Any] | None] = [
        {"workspace_root": str(tmp_path / "a"), "sandbox_root": str(tmp_path / "a/private"),
         "data_root": str(tmp_path / "state-a"), "shell_timeout": 7},
        {"workspace_root": str(tmp_path / "b"), "sandbox_root": str(tmp_path / "b/private"),
         "data_root": str(tmp_path / "state-b"), "shell_timeout": 11},
        {}, None,
    ]
    calls: list[tuple[Path, int]] = []

    def run_shell(args: Any, cwd: Path, timeout: int, allow: Any = ()) -> str:
        calls.append((cwd, timeout))
        return "executed"

    monkeypatch.setattr(coding, "_run_shell", run_shell)
    module = async_interaction if asynchronous else sync_interaction
    monkeypatch.setattr(module, "_prepare_shell", lambda command, shell, root, cwd: ([], 0, "test"))
    try:
        maps = []
        for session, config in zip(sessions, configs):
            session.user_input_provider = lambda *args: "y"
            maps.append({tool.tool_name: tool for tool in coding.get_tools(session, config)})
        first_config = configs[0]
        assert first_config is not None
        first_config["workspace_root"] = str(tmp_path / "b")
        for index in (0, 1, 2, 3, 0):
            expected = (tmp_path / ("a" if index == 0 else "b")).resolve() if index < 2 else coding.WORKSPACE_ROOT.resolve()
            tool = maps[index]["shell"]
            assert coding._tool_root(tool) == expected
            assert await _execute(tool, "test") == "executed"
            assert calls[-1] == (expected, (7, 11, 120, 120)[index])
            assert f"默认 {(7, 11, 120, 120)[index]}" in tool.params_dict["timeout"][1]
            definition = tool.get_tool_defined()["function"]["parameters"]
            assert definition["required"] == ["command"]
            assert definition["properties"]["timeout"]["default"] == (7, 11, 120, 120)[index]
            assert definition["properties"]["timeout"]["type"] == "integer"
            before = len(calls)
            assert "错误" in await _execute(tool, "test", timeout=0)
            assert len(calls) == before
            assert await _execute(tool, "test", timeout=19) == "executed"
            assert calls[-1][1] == 19
        for index in (0, 1, 0):
            root = tmp_path / ("a" if index == 0 else "b")
            (root / "private").mkdir(parents=True, exist_ok=True)
            (root / "source.txt").write_text(root.name, encoding="utf-8")
            assert root.name in await _execute(maps[index]["read_file"], "source.txt")
            sessions[index].user_input_provider = None
            assert "已写入" in await _execute(maps[index]["write_file"], "private/result.txt", root.name)
            assert "需要用户批准" in await _execute(maps[index]["write_file"], "outside.txt", "blocked")
            assert not (root / "outside.txt").exists()
            assert (root / "private/result.txt").read_text(encoding="utf-8") == root.name
        assert Path(get_plugin_state(sessions[0])["_data_root"]) == (tmp_path / "state-a").resolve()
        assert Path(get_plugin_state(sessions[1])["_data_root"]) == (tmp_path / "state-b").resolve()
    finally:
        for session in sessions:
            reset_plugin_state(session)

@pytest.mark.asyncio
@pytest.mark.parametrize("asynchronous", [False, True], ids=["sync", "async"])
async def test_real_install_state_roots_commands_plan_and_reinstall(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, asynchronous: bool,
) -> None:
    """真实安装保持状态根, 命令和处理器一致, 卸载后可切换数据目录"""
    from satrap.edictum import simple_session as ss_mod
    from satrap.edictum.plugin_config import PluginConfigManager

    monkeypatch.setattr(ss_mod, "PluginConfigManager", lambda: PluginConfigManager(tmp_path / "config"))
    base, llm_type = (AsyncSimpleSession, AsyncLLM) if asynchronous else (SimpleSession, LLM)
    session = base("same-id", MagicMock(spec=llm_type), db_path=str(tmp_path / "chat.db"))
    workspace = tmp_path / "workspace"
    (workspace / "private").mkdir(parents=True)
    config = {"workspace_root": str(workspace), "sandbox_root": str(workspace / "private"),
              "data_root": str(tmp_path / "state"), "shell_timeout": 7}

    async def install(cfg: dict[str, Any]) -> None:
        result = session.install_plugin(str(_coding_configuration_PLUGIN_DIR), config=cfg)
        if inspect.isawaitable(result):
            await result

    async def uninstall() -> None:
        result = session.uninstall_plugin("satrap_coding")
        assert (await result if inspect.isawaitable(result) else result) is True

    async def command(text: str) -> str:
        manager = session.command_handler if isinstance(session, AsyncSimpleSession) else session.cmd_handler
        result = manager.process_message(text)
        value, is_command = await result if inspect.isawaitable(result) else result
        assert is_command
        return str(value)

    await install(config)
    try:
        state = get_plugin_state(session)
        assert Path(state["_data_root"]) == (tmp_path / "state").resolve()
        tool = session.tools_manager.tools["write_file"]
        assert isinstance(tool, (coding.WriteFileTool, coding.AsyncWriteFileTool))
        assert tool.engine is state["engine"]
        assert "目标已设置" in await command("/goal shared-goal")
        assert (tmp_path / "state/goal.json").is_file()
        handler = session._handlers["satrap_coding.inject"]
        callback = handler.before_user_send
        assert callback is not None
        context = HandlerContext(HandlerConfig("message", None, "off", 1, "test"), "message")
        injected = callback("message", context)
        assert isinstance(injected, str) and "shared-goal" in injected
        assert "已切换" in await command("/approve mode full")
        assert (tmp_path / "state/permissions.json").is_file()
        assert "计划模式" in await command("/plan on")
        assert "计划模式" in await _execute(tool, "private/new.txt", "blocked")
        assert not (workspace / "private/new.txt").exists()
        assert "退出" in await command("/plan off")
        assert "已写入" in await _execute(tool, "outside.txt", "approved")
        assert (tmp_path / "state/approval_log.jsonl").is_file()
        await uninstall()
        await install({**config, "data_root": str(tmp_path / "new-state")})
        assert get_plugin_state(session) is not state
        replacement = session.tools_manager.tools["write_file"]
        assert isinstance(replacement, (coding.WriteFileTool, coding.AsyncWriteFileTool))
        assert not replacement.engine.plan_mode
        assert "shared-goal" not in await command("/goal status")
        assert not get_plugin_state(session)["goals"].to_context_block(session.session_id)
        assert Path(get_plugin_state(session)["_data_root"]) == (tmp_path / "new-state").resolve()
    finally:
        if session.list_plugins():
            await uninstall()

@pytest.mark.asyncio
@pytest.mark.parametrize("asynchronous", [False, True], ids=["sync", "async"])
@pytest.mark.parametrize("host_paths", [False, True], ids=["configured", "host"])
@pytest.mark.parametrize("order", [("tools", "commands", "handlers"), ("commands", "handlers", "tools"), ("handlers", "tools", "commands")])
async def test_factory_order_and_host_path_priority(
    tmp_path: Path, order: tuple[str, ...], asynchronous: bool, host_paths: bool,
) -> None:
    """工厂顺序不改变共享状态, 宿主路径优先级在两类会话中一致"""
    base, llm_type = (AsyncSimpleSession, AsyncLLM) if asynchronous else (SimpleSession, LLM)
    session = base("order", MagicMock(spec=llm_type), db_path=str(tmp_path / "chat.db"))
    if isinstance(session, AsyncSimpleSession):
        await session.initialize()
    if host_paths:
        setattr(session, "coding_workspace_root", str(tmp_path / "host-workspace"))
        setattr(session, "coding_sandbox_root", str(tmp_path / "host-sandbox"))
        setattr(session, "coding_cache_root", str(tmp_path / "host-cache"))
    config = {"workspace_root": str(tmp_path / "configured-workspace"), "sandbox_root": str(tmp_path / "configured-sandbox"), "data_root": str(tmp_path / "configured-data")}
    builders = {"tools": coding.get_tools, "commands": build_commands, "handlers": build_handlers}
    try:
        results = {name: builders[name](session, config) for name in order}
        state = get_plugin_state(session)
        expected_data = tmp_path / ("host-cache/satrap_coding" if host_paths else "configured-data")
        expected_workspace = tmp_path / ("host-workspace" if host_paths else "configured-workspace")
        assert Path(state["_data_root"]) == expected_data.resolve()
        tool = next(tool for tool in results["tools"] if tool.tool_name == "shell")
        assert coding._tool_root(tool) == expected_workspace.resolve()
        assert tool.engine is state["engine"]
        if host_paths:
            expected_workspace.mkdir()
            setattr(session, "coding_sandbox_root", str(expected_workspace))
            write_tool = next(tool for tool in results["tools"] if tool.tool_name == "write_file")
            assert "已写入" in await _execute(write_tool, "host.txt", "host")
            assert (expected_workspace / "host.txt").read_text(encoding="utf-8") == "host"
        sync_commands, async_commands = results["commands"]
        if asynchronous:
            await async_commands["plan"]("on")
        else:
            sync_commands["plan"]("on")
        assert tool.engine.plan_mode
        with pytest.raises(ValueError, match="卸载"):
            get_plugin_state(session, tmp_path / "other-data")
        assert tool.engine is get_plugin_state(session)["engine"]
    finally:
        reset_plugin_state(session)

@pytest.mark.asyncio
@pytest.mark.parametrize("asynchronous", [False, True], ids=["sync", "async"])
async def test_failed_install_clears_state_and_duplicate_install_preserves_it(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, asynchronous: bool,
) -> None:
    """失败安装释放状态, 重复安装不破坏已加载插件"""
    from satrap.edictum import simple_session as ss_mod
    from satrap.edictum.plugin_config import PluginConfigManager
    from satrap.expend.plugins.satrap_coding import state as state_mod

    monkeypatch.setattr(ss_mod, "PluginConfigManager", lambda: PluginConfigManager(tmp_path / "config"))
    base, llm_type = (AsyncSimpleSession, AsyncLLM) if asynchronous else (SimpleSession, LLM)
    session = base("retry", MagicMock(spec=llm_type), db_path=str(tmp_path / "chat.db"))
    manager = session.command_handler if isinstance(session, AsyncSimpleSession) else session.cmd_handler
    manager.register_command("goal", lambda *args: "existing")

    async def install(root: str) -> None:
        result = session.install_plugin(str(_coding_configuration_PLUGIN_DIR), config={"data_root": str(tmp_path / root)})
        if inspect.isawaitable(result):
            await result

    with pytest.raises(ValueError, match="命令.*冲突"):
        await install("failed-state")
    assert session.list_tools() == []
    assert session.list_plugins() == []
    assert id(session) not in state_mod._registry
    manager.unregister_command("goal")
    await install("retry-state")
    try:
        state = get_plugin_state(session)
        state["engine"].set_plan_mode(True)
        with pytest.raises(ValueError, match="已安装"):
            await install("different-state")
        assert get_plugin_state(session) is state
        assert state["engine"].plan_mode
        assert Path(state["_data_root"]) == (tmp_path / "retry-state").resolve()
    finally:
        result = session.uninstall_plugin("satrap_coding")
        if inspect.isawaitable(result):
            await result

@pytest.mark.asyncio
@pytest.mark.parametrize("asynchronous", [False, True], ids=["sync", "async"])
@pytest.mark.parametrize("tool_name,args", [
    ("write_file", ("file.txt", "changed")),
    ("edit_file", ("file.txt", "original", "changed")),
    ("search_replace", ("file.txt", [{"old": "original", "new": "changed"}])),
])
@pytest.mark.parametrize("change", ["workspace", "plan"])
async def test_file_approval_rechecks_execution_context(
    tmp_path: Path, asynchronous: bool, tool_name: str, args: tuple[Any, ...], change: str,
) -> None:
    """审批期间改绑工作区或进入计划模式时, 所有文件写工具取消执行"""
    base, llm_type = (AsyncSimpleSession, AsyncLLM) if asynchronous else (SimpleSession, LLM)
    session = base("approval", MagicMock(spec=llm_type), db_path=str(tmp_path / "chat.db"))
    if isinstance(session, AsyncSimpleSession):
        await session.initialize()
    root = tmp_path / "workspace"
    root.mkdir()
    target = root / "file.txt"
    target.write_text("original", encoding="utf-8")
    tools = {tool.tool_name: tool for tool in coding.get_tools(session, {"workspace_root": str(root)})}
    engine = get_plugin_state(session)["engine"]

    def approve(*args: Any) -> str:
        if change == "workspace":
            setattr(session, "coding_workspace_root", str(tmp_path / "other"))
        else:
            engine.set_plan_mode(True)
        return "y"

    session.user_input_provider = approve
    try:
        assert "执行已取消" in await _execute(tools[tool_name], *args)
        assert target.read_text(encoding="utf-8") == "original"
    finally:
        reset_plugin_state(session)

@pytest.mark.asyncio
@pytest.mark.parametrize("asynchronous", [False, True], ids=["sync", "async"])
@pytest.mark.parametrize("tool_name,args", [
    ("write_file", ("file.txt", "changed")),
    ("edit_file", ("file.txt", "original", "changed")),
    ("search_replace", ("file.txt", [{"old": "original", "new": "changed"}])),
    ("shell", ("echo x > f2.txt",)),
])
async def test_plan_mode_blocks_all_sandbox_file_writes(
    tmp_path: Path, asynchronous: bool, tool_name: str, args: tuple[Any, ...],
) -> None:
    """计划模式优先于沙箱免审批, 同步与异步文件/命令写入均被拒绝"""
    base, llm_type = (AsyncSimpleSession, AsyncLLM) if asynchronous else (SimpleSession, LLM)
    session = base("sandbox-plan", MagicMock(spec=llm_type), db_path=str(tmp_path / "chat.db"))
    if isinstance(session, AsyncSimpleSession):
        await session.initialize()
    target = tmp_path / "file.txt"
    target.write_text("original", encoding="utf-8")
    config = {"workspace_root": str(tmp_path), "sandbox_root": str(tmp_path)}
    tools = {tool.tool_name: tool for tool in coding.get_tools(session, config)}
    get_plugin_state(session)["engine"].set_plan_mode(True)
    try:
        assert "计划模式" in await _execute(tools[tool_name], *args)
        if tool_name != "shell":
            assert target.read_text(encoding="utf-8") == "original"
    finally:
        reset_plugin_state(session)

@pytest.mark.asyncio
@pytest.mark.parametrize("asynchronous", [False, True], ids=["sync", "async"])
async def test_host_sandbox_inside_system_data_remains_usable_and_scoped(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, asynchronous: bool,
) -> None:
    """平台默认数据目录内的当前会话沙箱可用, 其他运行数据仍受保护"""
    base, llm_type = (AsyncSimpleSession, AsyncLLM) if asynchronous else (SimpleSession, LLM)
    session = base("platform", MagicMock(spec=llm_type), db_path=str(tmp_path / "chat.db"))
    if isinstance(session, AsyncSimpleSession):
        await session.initialize()
    system_root = tmp_path / ".satrap"
    sandbox = system_root / "data/platforms/p/sessions/current/sandbox"
    sandbox.mkdir(parents=True)
    monkeypatch.setattr(coding, "_SYSTEM_PROTECTED_ROOT", system_root.resolve())
    setattr(session, "coding_workspace_root", str(sandbox))
    setattr(session, "coding_sandbox_root", str(sandbox))
    tools = {tool.tool_name: tool for tool in coding.get_tools(session)}
    try:
        assert "已写入" in await _execute(tools["write_file"], "normal.txt", "original")
        assert "original" in await _execute(tools["read_file"], "normal.txt")
        assert "已编辑" in await _execute(tools["edit_file"], "normal.txt", "original", "changed")
        assert "已批量替换" in await _execute(tools["search_replace"], "normal.txt", [{"old": "changed", "new": "updated"}])
        assert "normal.txt" in await _execute(tools["glob_files"], "*.txt")
        assert "updated" in await _execute(tools["grep_files"], "updated")
        for protected in (".env", ".satrap/private.txt", ".git/private.txt", "node_modules/private.txt"):
            assert "拒绝" in await _execute(tools["write_file"], protected, "blocked")
        other = system_root / "data/platforms/p/sessions/other/sandbox/private.txt"
        assert "越出工作区" in await _execute(tools["write_file"], str(other), "blocked")
        setattr(session, "coding_workspace_root", str(system_root))
        assert "拒绝" in await _execute(tools["write_file"], "credentials.txt", "blocked")
        assert not (system_root / "credentials.txt").exists()
        setattr(session, "coding_workspace_root", str(sandbox))
        delattr(session, "coding_sandbox_root")
        assert "拒绝" in await _execute(tools["read_file"], "normal.txt")
    finally:
        reset_plugin_state(session)

@pytest.mark.asyncio
@pytest.mark.parametrize("asynchronous", [False, True], ids=["sync", "async"])
@pytest.mark.parametrize("embedded", [False, True], ids=["ordinary", "embedded"])
@pytest.mark.parametrize("source", ["default", "configured", "host"])
@pytest.mark.parametrize("relation", ["same", "ancestor", "descendant"])
async def test_effective_sandbox_protection_matches_file_approval(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, asynchronous: bool,
    embedded: bool, source: str, relation: str,
) -> None:
    """实际沙箱的普通文件可用, 祖先工作区不会放开其他运行数据或敏感目录"""
    project = tmp_path / "project"
    system = project / ".satrap"
    default = system / "sandbox"
    configured = system / "coding/sandbox"
    sandbox = {"default": default, "configured": configured, "host": system / "data/sessions/current/sandbox"}[source]
    workspace = {"same": sandbox, "ancestor": project, "descendant": sandbox / "nested"}[relation]
    target_dir = workspace if relation == "descendant" else sandbox
    target_dir.mkdir(parents=True)
    monkeypatch.setattr(coding, "WORKSPACE_ROOT", project)
    monkeypatch.setattr(coding, "DEFAULT_SANDBOX_ROOT", default)
    monkeypatch.setattr(coding, "_SYSTEM_PROTECTED_ROOT", system)
    base, llm_type = (AsyncSimpleSession, AsyncLLM) if asynchronous else (SimpleSession, LLM)
    if embedded:
        session: Any = object.__new__(base)
        session.session_id = "embedded-sandbox"
        session._wf = SimpleNamespace(llm=None, tools_manager=None)
        session.session_ctx = SimpleNamespace(db_path=str(tmp_path / "chat.db"))
    else:
        session = base("ordinary-sandbox", MagicMock(spec=llm_type), db_path=str(tmp_path / "chat.db"))
        if isinstance(session, AsyncSimpleSession):
            await session.initialize()
    if source == "host":
        setattr(session, "coding_sandbox_root", str(sandbox))
    config = {"workspace_root": str(workspace), "protected_dirs": "private"}
    if source != "default":
        config["sandbox_root"] = str(configured)
    tools = {tool.tool_name: tool for tool in coding.get_tools(session, config)}
    target = target_dir / "normal.txt"
    relative = str(target.relative_to(workspace))
    try:
        assert "已写入" in await _execute(tools["write_file"], relative, "original")
        assert "original" in await _execute(tools["read_file"], relative)
        assert "已编辑" in await _execute(tools["edit_file"], relative, "original", "changed")
        assert "已批量替换" in await _execute(tools["search_replace"], relative, [{"old": "changed", "new": "updated"}])
        assert "normal.txt" in await _execute(tools["glob_files"], str(target_dir.relative_to(workspace) / "*.txt"))
        assert "updated" in await _execute(tools["grep_files"], "updated", str(target_dir.relative_to(workspace)))
        for protected in (".env", ".satrap/secret.txt", ".git/secret.txt", "node_modules/secret.txt", "private/secret.txt"):
            path = target_dir / protected
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_text("protected-marker", encoding="utf-8")
            name = str(path.relative_to(workspace))
            assert "拒绝" in await _execute(tools["write_file"], name, "blocked")
            assert "拒绝" in await _execute(tools["read_file"], name)
            assert path.read_text(encoding="utf-8") == "protected-marker"
        assert "protected-marker" not in await _execute(tools["grep_files"], "protected-marker")
        assert "secret.txt" not in await _execute(tools["glob_files"], "**/*.txt")
        get_plugin_state(session)["engine"].set_plan_mode(True)
        assert "计划模式" in await _execute(tools["write_file"], relative, "blocked")
        assert target.read_text(encoding="utf-8") == "updated"
        setattr(session, "coding_workspace_root", str(project))
        for sibling in (system / "credentials.txt", system / "data/sessions/other/sandbox/secret.txt", configured / "unused.txt" if source == "host" else system / "unused/secret.txt"):
            sibling.parent.mkdir(parents=True, exist_ok=True)
            sibling.write_text("outside-marker", encoding="utf-8")
            name = str(sibling.relative_to(project))
            assert "拒绝" in await _execute(tools["read_file"], name)
            assert "拒绝" in await _execute(tools["write_file"], name, "blocked")
            assert sibling.read_text(encoding="utf-8") == "outside-marker"
        assert "outside-marker" not in await _execute(tools["grep_files"], "outside-marker")
        for ancestor in (".git", "node_modules", "private"):
            blocked_sandbox = system / ancestor / "sandbox"
            blocked_sandbox.mkdir(parents=True)
            secret = blocked_sandbox / "secret.txt"
            secret.write_text("protected-ancestor", encoding="utf-8")
            setattr(session, "coding_sandbox_root", str(blocked_sandbox))
            assert "拒绝" in await _execute(tools["read_file"], str(secret.relative_to(project)))
    finally:
        reset_plugin_state(session)

@pytest.mark.parametrize("value", [0, -1, 3601, 1.5, True, "bad", float("inf"), float("nan")])
def test_factory_rejects_invalid_default_timeout(tmp_path: Path, value: object) -> None:
    """非法默认超时不能被截断或回退成可执行配置"""
    session = SimpleSession("invalid", MagicMock(spec=LLM), db_path=str(tmp_path / "chat.db"))
    try:
        with pytest.raises(ValueError, match="shell_timeout"):
            coding.get_tools(session, {"shell_timeout": value})
    finally:
        reset_plugin_state(session)

_satrap_coding_plugin_PLUGIN_DIR = Path(__file__).resolve().parents[2] / "satrap" / "expend" / "plugins" / "satrap_coding"

class _satrap_coding_plugin_FakeLLM(LLM):
    """记录调用参数的同步 fake LLM"""

    def __init__(self) -> None:
        self.calls: list[dict[str, Any]] = []

    def call(
        self, messages: list[dict[str, Any]], model: str | None = None,
        thinking: str = "off", temperature: float | None = None,
        top_p: float | None = None, max_tokens: int | None = None,
        tools: list[dict[str, Any]] | None = None,
        tool_choice: str = "auto", img_urls: list[str] | None = None,
    ) -> LLMCallResponse:
        self.calls.append({"messages": messages, "tools": tools})
        return LLMCallResponse(type="answer", content="回复")

    def stream_call(
        self, messages: list[dict[str, Any]], model: str | None = None,
        thinking: str = "off", temperature: float | None = None,
        top_p: float | None = None, max_tokens: int | None = None,
        tools: list[dict[str, Any]] | None = None,
        tool_choice: str = "auto", img_urls: list[str] | None = None,
    ) -> Iterator[LLMCallStreamEvent]:
        self.calls.append({"messages": messages, "tools": tools})
        yield LLMCallStreamEvent(kind="content_delta", delta="流")

@pytest.fixture
def _satrap_coding_plugin_session(tmp_path: Any, monkeypatch: Any) -> SimpleSession:
    """隔离插件数据目录的会话 + 安装 satrap_coding 插件 (经 config 指定工作区/数据目录)"""
    from satrap.edictum import simple_session as ss_mod
    from satrap.edictum.plugin_config import PluginConfigManager
    monkeypatch.setattr(ss_mod, "PluginConfigManager", lambda: PluginConfigManager(tmp_path / "cfg"))
    (tmp_path / "workspace").mkdir()
    s = SimpleSession(
        "conv-1", _satrap_coding_plugin_FakeLLM(), db_path=str(tmp_path / "chat.db"), enable_checkpoint=True,
    )
    s.install_plugin(str(_satrap_coding_plugin_PLUGIN_DIR), config={
        "workspace_root": str(tmp_path / "workspace"),
        "data_root": str(tmp_path / "coding"),
    })
    return s

def _user_text(s: SimpleSession, index: int = -1) -> str:
    """取第 index 次模型调用的最后一条用户消息"""
    calls = cast(list[dict[str, Any]], getattr(s.llm, "calls"))
    messages = calls[index]["messages"]
    return str([m for m in messages if m.get("role") == "user"][-1]["content"])

def test_plugin_install_full_capabilities(_satrap_coding_plugin_session: SimpleSession):
    """安装: 工具/命令/技能/处理器与能力描述一致 (search/memory 已移交 base_take)"""
    plugin = _satrap_coding_plugin_session.list_plugins()[0]
    assert plugin.name == "satrap_coding"
    assert len(plugin.tools) == 11
    assert set(plugin.commands) == {"goal", "plan", "approve"}
    assert set(plugin.skills) == {"goal", "plan"}
    assert set(plugin.handlers) == {"satrap_coding.inject"}
    assert len(_satrap_coding_plugin_session.list_tools()) == 11
    assert set(_satrap_coding_plugin_session.list_commands()) >= {"goal", "plan", "approve"}
    assert "goal" in _satrap_coding_plugin_session.list_skills()

    # meta.yaml 能力声明读入 plugin.capability_descriptions
    desc = plugin.capability_descriptions
    assert desc["tools"]["read_file"] and desc["tools"]["shell"]
    assert desc["skills"]["goal"] and desc["handlers"]["satrap_coding.inject"]
    assert desc["commands"]["approve"]
    assert desc.get("mcp", {}) == {}

    # list_capabilities 每项带 name / enabled / description
    caps = plugin.list_capabilities()
    for kind in ("tools", "skills", "handlers", "commands"):
        assert caps[kind], f"{kind} 应非空"
        for item in caps[kind]:
            assert "name" in item and "enabled" in item and "description" in item
    tools = {t["name"]: t for t in caps["tools"]}
    assert tools["read_file"]["description"] and tools["read_file"]["enabled"] is True

def test_goal_command_lifecycle_and_injection(_satrap_coding_plugin_session: SimpleSession):
    """/goal: 设置 -> 注入模型输入 -> status -> done/clear 后均不再注入"""
    result, is_cmd = _satrap_coding_plugin_session.cmd_handler.process_message("/goal 实现一个编码助手")
    assert is_cmd and "目标已设置" in str(result)

    _satrap_coding_plugin_session.run("帮我看看这个文件")
    text = _user_text(_satrap_coding_plugin_session)
    assert "实现一个编码助手" in text
    assert "【持续目标】" in text

    result, _ = _satrap_coding_plugin_session.cmd_handler.process_message("/goal status")
    assert "状态: active" in str(result)

    result, _ = _satrap_coding_plugin_session.cmd_handler.process_message("/goal done")
    assert "已标记完成" in str(result)

    _satrap_coding_plugin_session.run("再来一轮")
    assert "实现一个编码助手" not in _user_text(_satrap_coding_plugin_session)

    # clear 等价于终止注入
    _satrap_coding_plugin_session.cmd_handler.process_message("/goal 重构插件系统")
    _satrap_coding_plugin_session.run("第一轮")
    assert "重构插件系统" in _user_text(_satrap_coding_plugin_session)
    result, _ = _satrap_coding_plugin_session.cmd_handler.process_message("/goal clear")
    assert "已清除" in str(result)
    _satrap_coding_plugin_session.run("第二轮")
    assert "重构插件系统" not in _user_text(_satrap_coding_plugin_session)

def test_goal_todo(_satrap_coding_plugin_session: SimpleSession):
    """/goal todo: 子任务增删"""
    _satrap_coding_plugin_session.cmd_handler.process_message("/goal 写插件")
    result, _ = _satrap_coding_plugin_session.cmd_handler.process_message("/goal todo 写权限引擎")
    assert "子任务已添加" in str(result)
    _satrap_coding_plugin_session.cmd_handler.process_message("/goal todo 写测试")
    invalid, _ = _satrap_coding_plugin_session.cmd_handler.process_message("/goal todo-done 0")
    assert "序号无效" in str(invalid)
    result, _ = _satrap_coding_plugin_session.cmd_handler.process_message("/goal todo-done 1")
    assert "已标记完成" in str(result)
    result, _ = _satrap_coding_plugin_session.cmd_handler.process_message("/goal status")
    assert "1. [x] 写权限引擎" in str(result)
    assert "2. [ ] 写测试" in str(result)

def test_uninstall_resets_plugin_state(_satrap_coding_plugin_session: SimpleSession):
    """/plan on 后卸载: 工具/命令/处理器全部回收, 重装后 plan_mode 不残留"""
    _satrap_coding_plugin_session.cmd_handler.process_message("/plan on")

    assert _satrap_coding_plugin_session.uninstall_plugin("satrap_coding") is True
    assert _satrap_coding_plugin_session.list_tools() == []
    assert "goal" not in _satrap_coding_plugin_session.list_commands()
    assert "memory" not in _satrap_coding_plugin_session.list_commands()
    assert _satrap_coding_plugin_session.list_handlers() == []
    assert _satrap_coding_plugin_session.list_skills() == []
    assert _satrap_coding_plugin_session.list_plugins() == []

    _satrap_coding_plugin_session.install_plugin(str(_satrap_coding_plugin_PLUGIN_DIR))
    write_tool = _satrap_coding_plugin_session.tools_manager.tools["write_file"]
    # 重装后计划模式已重置: 写文件进入正常审批流而不再是计划模式拒绝
    target = tools_mod.WORKSPACE_ROOT / "y.txt"
    out = write_tool.execute(str(target), "hi")
    assert "需要用户批准" in out
    assert "拒绝: 计划模式" not in out

def test_approve_command(_satrap_coding_plugin_session: SimpleSession):
    """/approve: 策略切换与持久规则"""
    write_tool = _satrap_coding_plugin_session.tools_manager.tools["write_file"]
    target = tools_mod._tool_root(write_tool) / "x.txt"

    out = write_tool.execute(str(target), "hi")
    # user 模式无输入通道: 拒绝
    assert "需要用户批准" in out

    result, _ = _satrap_coding_plugin_session.cmd_handler.process_message("/approve mode full")
    assert "已切换: full" in str(result)
    assert "已写入" in write_tool.execute(str(target), "hi")

    result, _ = _satrap_coding_plugin_session.cmd_handler.process_message("/approve rule file_write 1")
    assert "持久规则已添加" in str(result)
    result, _ = _satrap_coding_plugin_session.cmd_handler.process_message("/approve rules")
    assert "file_write" in str(result)

def test_parse_capability_descriptions_unit():
    """parse_capability_descriptions: 字典解析 / 非 dict 跳过 / 未识别键忽略"""
    from satrap.edictum.plugin import parse_capability_descriptions

    meta: dict[str, Any] = {
        "name": "p",
        "tools": {"a": "描述a", "b": 2},
        "skills": {"s": "技能s"},
        "handlers": "not-a-dict",   # 非 dict -> 跳过
        "unknown": {"x": "y"},   # 未识别键 -> 忽略
    }
    desc = parse_capability_descriptions(meta)
    assert desc["tools"] == {"a": "描述a", "b": "2"}   # 值统一 str()
    assert desc["skills"] == {"s": "技能s"}
    assert "handlers" not in desc
    assert "unknown" not in desc

def test_undeclared_capability_warns_only(tmp_path: Any, monkeypatch: Any, caplog: Any):
    """声明了但扫描不到的能力仅警告, 不改变安装行为"""
    import logging
    from satrap.edictum.simple_session import _warn_undeclared_capabilities

    descriptions = {"tools": {"ghost": "不存在的工具", "real": "真实工具"}}
    scanned = {"tools": {"real": True}}
    with caplog.at_level(logging.WARNING):
        _warn_undeclared_capabilities("p", descriptions, scanned)
    assert any("ghost" in r.message for r in caplog.records)
    assert not any("real" in r.message and "未在插件中发现" in r.message for r in caplog.records)

_base_take_plugin_PLUGIN_DIR = Path(__file__).resolve().parents[2] / "satrap" / "expend" / "plugins" / "base_take"

CODING_PLUGIN_DIR = Path(__file__).resolve().parents[2] / "satrap" / "expend" / "plugins" / "satrap_coding"

class _ReportlabCanvas(Protocol):
    """reportlab Canvas 的最小调用面 (可选依赖, 无类型声明)"""

    def drawString(self, x: float, y: float, text: str) -> None: ...
    def save(self) -> None: ...

class _ReportlabCanvasModule(Protocol):
    """reportlab.pdfgen.canvas 模块的最小调用面"""

    Canvas: Callable[[str], _ReportlabCanvas]

class _base_take_plugin_FakeLLM(LLM):
    """最小同步 fake LLM"""

    def __init__(self) -> None:
        pass

    def call(
        self, messages: list[dict[str, Any]], model: str | None = None,
        thinking: str = "off", temperature: float | None = None,
        top_p: float | None = None, max_tokens: int | None = None,
        tools: list[dict[str, Any]] | None = None,
        tool_choice: str = "auto", img_urls: list[str] | None = None,
    ) -> LLMCallResponse:
        return LLMCallResponse(type="answer", content="回复")

    def stream_call(
        self, messages: list[dict[str, Any]], model: str | None = None,
        thinking: str = "off", temperature: float | None = None,
        top_p: float | None = None, max_tokens: int | None = None,
        tools: list[dict[str, Any]] | None = None,
        tool_choice: str = "auto", img_urls: list[str] | None = None,
    ) -> Iterator[LLMCallStreamEvent]:
        yield LLMCallStreamEvent(kind="content_delta", delta="流")

@pytest.fixture
def _base_take_plugin_session(tmp_path: Any, monkeypatch: Any) -> SimpleSession:
    """隔离数据目录的会话 + 安装 base_take 插件"""
    from satrap.edictum import simple_session as ss_mod
    from satrap.edictum.plugin_config import PluginConfigManager

    monkeypatch.setattr(ss_mod, "PluginConfigManager", lambda: PluginConfigManager(tmp_path / "cfg"))
    # 隔离: 插件配置目录 + 记忆 db + 沙箱根

    s = SimpleSession("conv-1", _base_take_plugin_FakeLLM(), db_path=str(tmp_path / "chat.db"))
    s.install_plugin(str(_base_take_plugin_PLUGIN_DIR), config={
        "sandbox_root": str(tmp_path / "sandbox"),
        "workspace_root": str(tmp_path / "workspace"),
    })
    (tmp_path / "workspace").mkdir(exist_ok=True)
    return s

def test_plugin_install_capabilities(_base_take_plugin_session: SimpleSession):
    """安装: 4 工具, 无记忆处理器和命令"""
    plugin = _base_take_plugin_session.list_plugins()[0]
    assert plugin.name == "base_take"
    assert set(plugin.tools) == {
        "search", "fetch_page", "code_sandbox", "read_document",
    }
    assert not plugin.handlers
    assert not plugin.commands
    assert len(_base_take_plugin_session.list_tools()) == 4

def test_config_schema_on_plugin(_base_take_plugin_session: SimpleSession):
    """config_schema 随插件返回 (供前端渲染表单)"""
    plugin = _base_take_plugin_session.list_plugins()[0]
    assert plugin.config_schema["sandbox_root"]["type"] == "path"
    assert plugin.config_schema["search_timeout"]["default"] == 10
    assert "memory_scope" not in plugin.config_schema
    assert "memory_mode" not in plugin.config_schema


def test_base_take_sandbox_whitelist_and_description(_base_take_plugin_session: SimpleSession, tmp_path: Path) -> None:
    """工厂冻结生效白名单, 描述一致且不触发用户输入通道"""
    from satrap.expend.plugins.base_take.tools import get_tools
    from satrap.expend.tools.sandbox_tools import CodeSandboxTool

    def never_approve(*args: object) -> str:
        pytest.fail("code_sandbox 不应调用用户审批通道")

    _base_take_plugin_session.user_input_provider = never_approve
    config = {"sandbox_root": str(tmp_path / "configured"), "sandbox_allowed_modules": "math"}
    first = next(tool for tool in get_tools(_base_take_plugin_session, config) if isinstance(tool, CodeSandboxTool))
    config["sandbox_allowed_modules"] = ""
    second = next(tool for tool in get_tools(_base_take_plugin_session, config) if isinstance(tool, CodeSandboxTool))
    assert first.sandbox.modules == ("math",)
    assert second.sandbox.modules == ()
    assert "允许导入: math." in first.description and "允许导入: (无)" in second.description
    assert first.execute("run", "import math\nprint(math.sqrt(9))")["stdout"] == "3.0\n"
    assert second.execute("run", "import math")["error_code"] == "module_not_allowed"

def test_extract_text_txt(tmp_path: Path):
    """纯文本直接读取"""
    f = tmp_path / "a.txt"
    f.write_text("你好世界", encoding="utf-8")
    assert extract_text(f) == "你好世界"

def test_extract_text_xlsx(tmp_path: Path):
    """xlsx 解析为 TSV 文本"""
    from openpyxl import Workbook
    wb: Any = Workbook()   # openpyxl 无类型声明
    ws: Any = wb.active
    assert ws is not None
    ws.title = "数据"
    ws.append(["姓名", "年龄"])
    ws.append(["张三", 30])
    f = tmp_path / "t.xlsx"
    wb.save(str(f))
    wb.close()
    text = extract_text(f)
    assert "Sheet: 数据" in text
    assert "姓名" in text and "张三" in text

def test_extract_text_docx(tmp_path: Path):
    """docx 解析段落为纯文本"""
    import docx
    doc = docx.Document()
    doc.add_paragraph("第一段内容")
    doc.add_paragraph("第二段内容")
    f = tmp_path / "t.docx"
    doc.save(str(f))
    text = extract_text(f)
    assert "第一段内容" in text and "第二段内容" in text

def test_extract_text_pdf(tmp_path: Path):
    """pdf 解析 (用 pdfplumber 生成最小 pdf 较复杂, 改用 reportlab 若可用否则跳过)"""
    # reportlab 为可选依赖, 未安装时 importorskip 跳过; 调用面以下方 Protocol 静态声明
    canvas_mod = cast(_ReportlabCanvasModule, pytest.importorskip("reportlab.pdfgen.canvas"))
    f = tmp_path / "t.pdf"
    c = canvas_mod.Canvas(str(f))
    c.drawString(100, 750, "Hello PDF")
    c.save()
    text = extract_text(f)
    assert "Hello PDF" in text

def test_pdfminer_fontbbox_warning_filtered():
    """缺 FontBBox 的良性警告被过滤, 其他 pdfminer 警告不受影响, 挂载幂等"""
    import logging

    from satrap.core.utils.documents import _mute_pdfminer_fontbbox_warning

    class _Capture(logging.Handler):
        def __init__(self) -> None:
            super().__init__()
            self.messages: list[str] = []

        def emit(self, record: logging.LogRecord) -> None:
            self.messages.append(record.getMessage())

    logger = logging.getLogger("pdfminer.pdffont")
    _mute_pdfminer_fontbbox_warning()
    _mute_pdfminer_fontbbox_warning()   # 重复调用不叠加
    filters = [f for f in logger.filters if f.__class__.__name__ == "_FontBBoxWarningFilter"]
    assert len(filters) == 1

    capture = _Capture()
    logger.addHandler(capture)
    try:
        logger.warning("Could not get FontBBox from font descriptor because %r cannot be parsed as 4 floats", None)
        logger.warning("pdfminer 其他警告")
        assert capture.messages == ["pdfminer 其他警告"]
    finally:
        logger.removeHandler(capture)

@pytest.mark.parametrize("filename, content, match", [
    ("t.bin", b"\x00\x01", "不支持"),
    ("ghost.txt", None, "不存在"),
])
def test_extract_text_rejects_unsupported_or_missing(tmp_path: Path, filename: str, content: bytes | None, match: str):
    """不支持的类型或文件不存在时返回明确错误"""
    path = tmp_path / filename
    if content is not None:
        path.write_bytes(content)
    with pytest.raises(ValueError, match=match):
        extract_text(path)

def test_extract_text_enforces_file_and_output_limits(tmp_path: Path):
    """文本读取在解析前限制文件大小, 并在解析中限制输出长度"""
    path = tmp_path / "large.txt"
    path.write_text("abcdefghij", encoding="utf-8")
    assert extract_text(path, max_length=4) == "abcd"
    with pytest.raises(ValueError, match="文件过大"):
        extract_text(path, max_file_size=4)

def test_extract_text_rejects_extreme_zip_expansion(tmp_path: Path):
    """OOXML 在交给解析器前拒绝异常展开大小"""
    path = tmp_path / "bomb.docx"
    with ZipFile(path, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", "x" * 10_000)
    with pytest.raises(ValueError, match="展开后过大"):
        extract_text(path, max_expanded_size=100)

def test_read_document_tool(_base_take_plugin_session: SimpleSession, tmp_path: Path):
    """read_document 工具: 工作区内解析 + 越界拒绝"""
    workspace = tmp_path / "workspace"
    (workspace / "note.txt").write_text("工作区笔记", encoding="utf-8")
    tool = _base_take_plugin_session._wf.tools_manager.tools["read_document"]
    assert "工作区笔记" in tool.execute(path="note.txt")
    assert tool.execute(path="note.txt", max_length="bad") == "错误: max_length 必须是整数"
    # 越界拒绝
    result = tool.execute(path="../outside.txt")
    assert "越出工作区" in result or "错误" in result
