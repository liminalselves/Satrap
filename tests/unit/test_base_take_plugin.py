"""base_take 插件测试: 安装 / 工具组成 / 配置注入 / 文档解析 / 记忆注入"""
from __future__ import annotations

from pathlib import Path
from typing import Any, Iterator

import pytest

from satrap.core.APICall.LLMCall import LLM
from satrap.core.type import LLMCallResponse, LLMCallStreamEvent
from satrap.edictum import SimpleSession
from satrap.expend.plugins.base_take.core.docread import extract_text

PLUGIN_DIR = Path(__file__).resolve().parents[2] / "satrap" / "expend" / "plugins" / "base_take"
CODING_PLUGIN_DIR = Path(__file__).resolve().parents[2] / "satrap" / "expend" / "plugins" / "satrap_coding"


class _FakeLLM(LLM):
    """最小同步 fake LLM"""

    def __init__(self) -> None:
        pass

    def call(
        self, messages: list[dict[str, Any]], model: str | None = None,
        thinking: str = "off", temperature: float | None = None,
        top_p: float | None = None, max_tokens: int | None = None,
        tools: list[dict[str, Any]] | None = None,
        tool_choice: str = "auto", img_urls: list[str] | None = None,
    ) -> LLMCallResponse:
        return LLMCallResponse(type="answer", content="回复")

    def stream_call(
        self, messages: list[dict[str, Any]], model: str | None = None,
        thinking: str = "off", temperature: float | None = None,
        top_p: float | None = None, max_tokens: int | None = None,
        tools: list[dict[str, Any]] | None = None,
        tool_choice: str = "auto", img_urls: list[str] | None = None,
    ) -> Iterator[LLMCallStreamEvent]:
        yield LLMCallStreamEvent(kind="content_delta", delta="流")


@pytest.fixture
def session(tmp_path: Any, monkeypatch: Any) -> SimpleSession:
    """隔离数据目录的会话 + 安装 base_take 插件"""
    from satrap.edictum import simple_session as ss_mod
    from satrap.edictum.plugin_config import PluginConfigManager
    from satrap.expend.plugins.base_take import state as state_mod
    from satrap.expend.tools import memory_store as ms_mod

    # 隔离: 插件配置目录 + 记忆 db + 沙箱根
    monkeypatch.setattr(ss_mod, "PluginConfigManager", lambda: PluginConfigManager(tmp_path / "cfg"))
    monkeypatch.setattr(ms_mod, "DEFAULT_MEMORY_DB", tmp_path / "memory.db")
    monkeypatch.setattr(state_mod, "DEFAULT_MEMORY_DB", tmp_path / "memory.db")

    s = SimpleSession("conv-1", _FakeLLM(), db_path=str(tmp_path / "chat.db"))
    s.install_plugin(str(PLUGIN_DIR), config={
        "sandbox_root": str(tmp_path / "sandbox"),
        "workspace_root": str(tmp_path / "workspace"),
    })
    (tmp_path / "workspace").mkdir(exist_ok=True)
    return s


def test_plugin_install_capabilities(session: SimpleSession):
    """安装: 8 工具 + 1 处理器 + 1 命令"""
    plugin = session.list_plugins()[0]
    assert plugin.name == "base_take"
    assert set(plugin.tools) == {
        "search", "fetch_page", "code_sandbox", "read_document",
        "add_memory", "update_memory", "delete_memory", "list_memories",
    }
    assert set(plugin.handlers) == {"base_take.memory_inject"}
    assert set(plugin.commands) == {"memory"}
    assert len(session.list_tools()) == 8


def test_config_schema_on_plugin(session: SimpleSession):
    """config_schema 随插件返回 (供前端渲染表单)"""
    plugin = session.list_plugins()[0]
    assert plugin.config_schema["sandbox_root"]["type"] == "path"
    assert plugin.config_schema["search_timeout"]["default"] == 10
    assert plugin.config_schema["memory_scope"]["default"] == "web_chat"
    assert plugin.config_schema["memory_mode"]["default"] == "full"
    assert plugin.config_schema["memory_mode"]["options"] == ["disabled", "base", "full"]


def test_memory_mode_config_readonly(tmp_path: Any, monkeypatch: Any):
    """memory_mode=base: 写工具被拒绝; disabled: 注入块为空"""
    from satrap.edictum import simple_session as ss_mod
    from satrap.edictum.plugin_config import PluginConfigManager
    from satrap.expend.plugins.base_take import state as state_mod
    from satrap.expend.tools import memory_store as ms_mod

    monkeypatch.setattr(ss_mod, "PluginConfigManager", lambda: PluginConfigManager(tmp_path / "cfg"))
    monkeypatch.setattr(ms_mod, "DEFAULT_MEMORY_DB", tmp_path / "memory.db")
    monkeypatch.setattr(state_mod, "DEFAULT_MEMORY_DB", tmp_path / "memory.db")

    s = SimpleSession("conv-mode", _FakeLLM(), db_path=str(tmp_path / "chat.db"))
    s.install_plugin(str(PLUGIN_DIR), config={
        "sandbox_root": str(tmp_path / "sandbox"),
        "workspace_root": str(tmp_path / "workspace"),
        "memory_mode": "base",
    })

    from satrap.expend.plugins.base_take.state import get_plugin_state
    store = get_plugin_state(s)["store"]
    assert store.mode == "base"
    assert store.can_write() is False

    # base 只读: 写工具返回"只读"提示
    tm = s._wf.tools_manager
    add = tm.tools["add_memory"]
    assert "只读" in add.execute(title="偏好", content="x")

    # disabled 模式: 不注入 (注入块为空), 写拒绝文案按实际模式生成
    store.set_mode("disabled")
    assert store.to_context_block() == ""
    assert "已禁用" in add.execute(title="偏好", content="x")
    assert "disabled" in add.execute(title="偏好", content="x")


def test_memory_tools_lifecycle(session: SimpleSession):
    """memory 工具: 添加 -> 列表 -> 更新 -> 删除"""
    tm = session._wf.tools_manager
    add = tm.tools["add_memory"]
    result = add.execute(title="偏好", content="喜欢简洁回答", tags=["style"], importance=3)
    assert "已添加" in result

    list_tool = tm.tools["list_memories"]
    listed = list_tool.execute()
    assert "偏好" in listed and "喜欢简洁回答" in listed

    # 取记忆 ID 前缀
    from satrap.expend.plugins.base_take.state import get_plugin_state
    store = get_plugin_state(session)["store"]
    mem_id = store.list_all()[0]["id"]

    update = tm.tools["update_memory"]
    assert "已更新" in update.execute(memory_id=mem_id, content="喜欢极简回答")

    delete = tm.tools["delete_memory"]
    assert "已删除" in delete.execute(memory_id=mem_id)
    assert "没有" in list_tool.execute()


def test_memory_command_lifecycle(session: SimpleSession):
    """/memory 命令: 注册 -> list/add/del/mode 全流程"""
    cmd = session.cmd_handler.commands["memory"]

    assert "没有长期记忆" in cmd()
    assert "已添加" in cmd("add", "偏好", "喜欢简洁回答")
    listed = cmd()
    assert "偏好" in listed and "喜欢简洁回答" in listed

    mem_id = listed.split("- ")[1].split(" ")[0]
    assert "已删除" in cmd("del", mem_id[:8])  # 支持 ID 前缀

    assert "已切换: base" in cmd("mode", "base")
    assert "只读模式" in cmd("add", "x", "y")
    assert "只读模式" in cmd("clear")
    assert "已切换: full" in cmd("mode", "full")
    assert "用法" in cmd("unknown")


def test_memory_inject_handler(session: SimpleSession):
    """记忆注入 handler: 添加记忆后注入到用户输入"""
    from satrap.expend.plugins.base_take.state import get_plugin_state
    state = get_plugin_state(session)
    store = state["store"]
    store.add("项目约定", "回复用中文", importance=5)

    # 直接用注入器 (handler 已把它挂在 state 上)
    injector = state["_injector"]
    injected = injector.inject("你好")
    assert "长期记忆" in injected
    assert "项目约定" in injected
    assert "回复用中文" in injected


def test_memory_writes_not_blocked_by_plan_mode(tmp_path: Any, monkeypatch: Any):
    """计划模式只限制工作区写操作 (文件/shell/沙箱); 记忆是元信息, 增删改不受拦截 (有意设计)"""
    from satrap.edictum import simple_session as ss_mod
    from satrap.edictum.plugin_config import PluginConfigManager
    from satrap.expend.plugins.base_take import state as state_mod
    from satrap.expend.tools import memory_store as ms_mod

    monkeypatch.setattr(ss_mod, "PluginConfigManager", lambda: PluginConfigManager(tmp_path / "cfg"))
    monkeypatch.setattr(ms_mod, "DEFAULT_MEMORY_DB", tmp_path / "memory.db")
    monkeypatch.setattr(state_mod, "DEFAULT_MEMORY_DB", tmp_path / "memory.db")
    (tmp_path / "workspace").mkdir()

    s = SimpleSession("conv-plan", _FakeLLM(), db_path=str(tmp_path / "chat.db"))
    s.install_plugin(str(CODING_PLUGIN_DIR), config={
        "workspace_root": str(tmp_path / "workspace"),
        "data_root": str(tmp_path / "coding"),
    })
    s.install_plugin(str(PLUGIN_DIR), config={
        "sandbox_root": str(tmp_path / "sandbox"),
        "workspace_root": str(tmp_path / "workspace"),
    })

    assert "已进入计划模式" in s.cmd_handler.commands["plan"]("on")

    # 计划模式下记忆写不被拦截: 工具与命令均可写
    add = s._wf.tools_manager.tools["add_memory"]
    assert "已添加" in add.execute(title="决策", content="用 SQLite", tags=[], importance=1)
    assert "已添加" in s.cmd_handler.commands["memory"]("add", "偏好", "中文回复")


# ================= 文档解析 =================

def test_extract_text_txt(tmp_path: Path):
    """纯文本直接读取"""
    f = tmp_path / "a.txt"
    f.write_text("你好世界", encoding="utf-8")
    assert extract_text(f) == "你好世界"


def test_extract_text_xlsx(tmp_path: Path):
    """xlsx 解析为 TSV 文本"""
    from openpyxl import Workbook
    wb: Any = Workbook()  # openpyxl 无类型声明
    ws: Any = wb.active
    assert ws is not None
    ws.title = "数据"
    ws.append(["姓名", "年龄"])
    ws.append(["张三", 30])
    f = tmp_path / "t.xlsx"
    wb.save(str(f))
    wb.close()
    text = extract_text(f)
    assert "Sheet: 数据" in text
    assert "姓名" in text and "张三" in text


def test_extract_text_docx(tmp_path: Path):
    """docx 解析段落为纯文本"""
    import docx
    doc = docx.Document()
    doc.add_paragraph("第一段内容")
    doc.add_paragraph("第二段内容")
    f = tmp_path / "t.docx"
    doc.save(str(f))
    text = extract_text(f)
    assert "第一段内容" in text and "第二段内容" in text


def test_extract_text_pdf(tmp_path: Path):
    """pdf 解析 (用 pdfplumber 生成最小 pdf 较复杂, 改用 reportlab 若可用否则跳过)"""
    pytest.importorskip("reportlab")
    from reportlab.pdfgen import canvas  # type: ignore[reportMissingModuleSource] 可选依赖, 未安装时上面 importorskip 跳过
    f = tmp_path / "t.pdf"
    c = canvas.Canvas(str(f))
    c.drawString(100, 750, "Hello PDF")
    c.save()
    text = extract_text(f)
    assert "Hello PDF" in text


def test_pdfminer_fontbbox_warning_filtered():
    """缺 FontBBox 的良性警告被过滤, 其他 pdfminer 警告不受影响, 挂载幂等"""
    import logging

    from satrap.expend.plugins.base_take.core.docread import _mute_pdfminer_fontbbox_warning

    class _Capture(logging.Handler):
        def __init__(self) -> None:
            super().__init__()
            self.messages: list[str] = []

        def emit(self, record: logging.LogRecord) -> None:
            self.messages.append(record.getMessage())

    logger = logging.getLogger("pdfminer.pdffont")
    _mute_pdfminer_fontbbox_warning()
    _mute_pdfminer_fontbbox_warning()  # 重复调用不叠加
    filters = [f for f in logger.filters if f.__class__.__name__ == "_FontBBoxWarningFilter"]
    assert len(filters) == 1

    capture = _Capture()
    logger.addHandler(capture)
    try:
        logger.warning("Could not get FontBBox from font descriptor because %r cannot be parsed as 4 floats", None)
        logger.warning("pdfminer 其他警告")
        assert capture.messages == ["pdfminer 其他警告"]
    finally:
        logger.removeHandler(capture)


def test_extract_text_unsupported(tmp_path: Path):
    """不支持的类型返回明确错误"""
    f = tmp_path / "t.bin"
    f.write_bytes(b"\x00\x01")
    with pytest.raises(ValueError, match="不支持"):
        extract_text(f)


def test_extract_text_missing(tmp_path: Path):
    """文件不存在返回明确错误"""
    with pytest.raises(ValueError, match="不存在"):
        extract_text(tmp_path / "ghost.txt")


def test_read_document_tool(session: SimpleSession, tmp_path: Path):
    """read_document 工具: 工作区内解析 + 越界拒绝"""
    workspace = tmp_path / "workspace"
    (workspace / "note.txt").write_text("工作区笔记", encoding="utf-8")
    tool = session._wf.tools_manager.tools["read_document"]
    assert "工作区笔记" in tool.execute(path="note.txt")
    # 越界拒绝
    result = tool.execute(path="../outside.txt")
    assert "越出工作区" in result or "错误" in result
