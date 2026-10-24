"""业务链路场景与集中边界契约"""
from __future__ import annotations

from dataclasses import replace
from satrap.core.plugin_interop.environment import bind_environment, environment_snapshot
import json
from satrap.core.approval.contracts import ApprovalRequest, ApprovalResult

from pathlib import Path
import pytest
import sys
from satrap.expend.tools.sandbox_tools import AsyncCodeSandboxTool, CodeSandboxTool
from satrap.core.utils.sandbox import CodeSandbox
import shutil
from types import SimpleNamespace
import os
from typing import Any, cast
from satrap.core.approval.policy import PermissionEngine, RiskLevel
from satrap.expend.plugins.satrap_coding import tools as coding
from satrap.core.utils.proc_env import sanitized_child_env

@pytest.fixture
def sandbox(tmp_path: Path) -> CodeSandbox:
    """创建隔离的代码沙箱"""
    return CodeSandbox(str(tmp_path / "sandbox"), sys.executable)

@pytest.mark.parametrize("path", ["../outside.py", "nested/../../outside.py"])
def test_path_traversal_is_rejected_for_all_file_operations(
    sandbox: CodeSandbox,
    tmp_path: Path,
    path: str,
) -> None:
    """保存, 读取, 执行和删除均在越界时直接拒绝"""
    with pytest.raises(ValueError, match="越出沙箱范围"):
        sandbox.save_to_file("print('x')", path)
    assert sandbox.read_file(path)["returncode"] == -12
    with pytest.raises(ValueError, match="越出沙箱范围"):
        sandbox.run_file(path)
    with pytest.raises(ValueError, match="越出沙箱范围"):
        sandbox.delete_file(path)
    with pytest.raises(ValueError, match="越出沙箱范围"):
        sandbox.delete_directory(path)
    assert not (tmp_path / "outside.py").exists()

def test_absolute_outside_path_is_rejected(sandbox: CodeSandbox, tmp_path: Path) -> None:
    """绝对路径不能绕过沙箱根目录"""
    outside = tmp_path / "outside.py"
    with pytest.raises(ValueError, match="越出沙箱范围"):
        sandbox.save_to_file("print('x')", str(outside))
    with pytest.raises(ValueError, match="越出沙箱范围"):
        sandbox.list_files(str(tmp_path))

@pytest.mark.parametrize("path", ["", "."])
def test_delete_sandbox_root_is_rejected(sandbox: CodeSandbox, path: str) -> None:
    """删除目录操作不得删除沙箱根目录"""
    marker = Path(sandbox.sandbox_path) / "marker.txt"
    sandbox.save_to_file("keep", "marker.txt")
    with pytest.raises(PermissionError, match="沙箱根目录"):
        sandbox.delete_directory(path)
    assert marker.read_text(encoding="utf-8") == "keep"

def test_internal_file_lifecycle(sandbox: CodeSandbox) -> None:
    """合法的沙箱内文件操作保持可用"""
    sandbox.save_to_file("print('ok')", "nested/main.py")
    assert sandbox.read_file("nested/main.py")["content"] == "print('ok')"
    assert sandbox.list_files("nested") == [str(Path("nested") / "main.py")]
    sandbox.delete_file("nested/main.py")
    sandbox.delete_directory("nested")
    assert not Path(sandbox.sandbox_path, "nested").exists()

def test_code_execution_timeout(tmp_path: Path) -> None:
    """死循环代码在配置的截止时间后返回超时错误"""
    limited = CodeSandbox(str(tmp_path / "sandbox"), sys.executable, execution_timeout=0.1)
    result = limited.run("while True: pass")
    assert result["returncode"] == -4
    assert "超时" in result["stderr"]

@pytest.mark.parametrize("asynchronous", [False, True], ids=["sync", "async"])
async def test_sandbox_worker_lifecycle_without_approval(sandbox: CodeSandbox, asynchronous: bool) -> None:
    """七操作共用真实新进程, 无授权通道也能执行, 文件保留但变量不保留"""
    tool = AsyncCodeSandboxTool(sandbox) if asynchronous else CodeSandboxTool(sandbox)

    async def execute(operation, **kwargs):
        return await tool.execute(operation, **kwargs) if asynchronous else tool.execute(operation, **kwargs)

    source = "import math\nanswer=math.sqrt(81)\nprint('你好', answer)"
    calls = [
        await execute("save", code=source, path="nested/main.py"),
        await execute("read", path="nested/main.py"),
        await execute("run_file", path="nested/main.py"),
        await execute("run", code="print('fresh')"),
        await execute("list"),
        await execute("delete", path="nested/main.py"),
        await execute("delete_dir", path="nested"),
    ]
    assert all(r["returncode"] == 0 for r in calls), calls
    pids = [r["worker_pid"] for r in calls]
    assert len(set(pids)) == len(pids) and os.getpid() not in pids
    assert calls[1]["content"] == source
    assert calls[2]["stdout"] == "你好 9.0\n"
    assert calls[4]["files"] == [str(Path("nested") / "main.py")]
    missing = await execute("run", code="print(answer)")
    assert missing["error_code"] == "execution_failed" and "NameError" in missing["error"]
    assert not Path(sandbox.sandbox_path, "nested").exists()


def test_sandbox_analysis_and_document_artifacts(sandbox: CodeSandbox) -> None:
    """真实库完成数组/表格/绘图/Word/PPTX/XLSX/PDF 业务闭环"""
    source = """
import time
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Pt
from pptx import Presentation
from pptx.util import Inches
import openpyxl
from reportlab.pdfgen import canvas
from pypdf import PdfReader, PdfWriter
import pdfplumber

start = time.perf_counter()
a = np.arange(6).reshape(2,3)
a[0,0] = 2
np.save("array.npy", a)
print(np.load("array.npy").sum())
df = pd.DataFrame({"kind":["a","a","b"], "n":[1,2,3]})
df.loc[0,"n"] = 2
df.to_csv("data.csv", index=False)
df.to_json("data.json")
df.groupby("kind").sum().to_excel("data.xlsx")
print(pd.read_excel("data.xlsx")["n"].sum())
fig, ax = plt.subplots()
ax.plot([1,2,3], [3,2,4], label="Example")
ax.set_title("图表")
ax.legend()
fig.savefig("chart.png")
fig.savefig("chart.svg")
fig.savefig("chart.pdf")
plt.close(fig)
doc=Document()
doc.add_paragraph("Hello document").runs[0].font.size=Pt(12)
doc.add_table(rows=1,cols=2).cell(0,0).text="表格"
doc.add_picture("chart.png")
doc.save("report.docx")
print(len(Document("report.docx").paragraphs))
prs=Presentation()
slide=prs.slides.add_slide(prs.slide_layouts[6])
slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text="Hello slide"
slide.shapes.add_picture("chart.png", Inches(1), Inches(2), width=Inches(3))
prs.save("report.pptx")
print(len(Presentation("report.pptx").slides))
c=canvas.Canvas("report.pdf")
c.drawString(50,700,"Hello PDF")
c.save()
r=PdfReader("report.pdf")
w=PdfWriter()
w.add_page(r.pages[0])
w.write("copy.pdf")
with pdfplumber.open("copy.pdf") as pdf:
    print(pdf.pages[0].extract_text())
print(time.perf_counter() >= start)
"""
    result = sandbox.run(source)
    assert result["returncode"] == 0, result
    assert "Hello PDF" in result["stdout"] and "True" in result["stdout"]
    files = set(sandbox.list_files())
    assert {"chart.png","chart.svg","chart.pdf","array.npy","data.csv","data.json","data.xlsx","report.docx","report.pptx","report.pdf","copy.pdf"} <= files
    assert Path(sandbox.sandbox_path, "chart.png").read_bytes().startswith(b"\x89PNG")
    assert sandbox.read_file("chart.png").get("error")  # 二进制产物不伪装为文本附件
    assert result["elapsed_seconds"] < sandbox.execution_timeout


def test_sandbox_import_api_and_path_policy_flows(sandbox: CodeSandbox, tmp_path: Path) -> None:
    """导入, 普通别名和已登记库次级文件入口共用结构化拒绝"""
    rejected = sandbox.run("print('must not run')\nif False:\n import socket")
    assert rejected["error_code"] == "module_not_allowed" and rejected["stdout"] == ""
    sandbox.save_to_file("import subprocess", "forbidden.py")
    assert sandbox.run_file("forbidden.py")["error_code"] == "module_not_allowed"
    outside = tmp_path / "outside.txt"
    outside.write_text("private-sentinel", encoding="utf-8")
    cases = [
        (f"open({str(outside)!r}).read()", "path_outside_sandbox"),
        (f"import pandas as pd\npd.read_csv({str(outside)!r})", "path_outside_sandbox"),
        ("import pandas as pd\npd.read_csv('http://localhost/secret')", "path_outside_sandbox"),
        ("import pandas as pd\nf=pd.DataFrame({'a':[1]}).query\nf('a==1')", "api_not_allowed"),
        ("import numpy as np\nnp.save('unsafe.npy',np.array([1]),allow_pickle=True)", "argument_not_allowed"),
        ("import numpy as np\nnp.ctypeslib", "api_not_allowed"),
        ("import matplotlib.pyplot as plt\nplt.rcParams['text.usetex']=True", "argument_not_allowed"),
        ("import matplotlib.pyplot as plt\nplt.plot([1,2])\nplt.savefig('bad.png',backend='pgf')", "argument_not_allowed"),
        ("from reportlab.platypus import DocExec", "api_not_allowed"),
        ("import time\ntime.clock_settime", "api_not_allowed"),
    ]
    for source, expected in cases:
        result = sandbox.run(source)
        assert result.get("error_code") == expected, (source, result)
    caught = sandbox.run("import numpy as np\ntry:\n np.ctypeslib\nexcept Exception:\n pass\nprint('continued')")
    assert caught["error_code"] == "api_not_allowed" and caught["stdout"] == "continued\n"
    assert outside.read_text(encoding="utf-8") == "private-sentinel"

    text = sandbox.run("print('中'*40000)")
    assert text["stdout_truncated"] and len(text["stdout"]) == 20_000
    unsupported = sandbox.run("input('never wait')")
    assert unsupported["returncode"] != 0 and unsupported["error_code"] != "timeout"
    secret_script = sandbox.run("import numpy as np\nnp.savez('object.npz', data=np.array([{'x':1}], dtype='object'))")
    assert secret_script["error_code"] == "argument_not_allowed"
    if os.name == "nt":
        assert sandbox.execute("save", "bad", "NUL")["error_code"] == "path_outside_sandbox"


async def test_sandbox_cancellation_and_configuration(tmp_path: Path, monkeypatch) -> None:
    """取消等待会回收真实 worker, 空/错误配置不能扩大能力"""
    import asyncio
    from satrap.core.utils.sandbox import runner
    processes = []
    original = runner.subprocess.Popen

    def spawn(*args, **kwargs):
        process = original(*args, **kwargs)
        processes.append(process)
        return process

    monkeypatch.setattr(runner.subprocess, "Popen", spawn)
    box = CodeSandbox(str(tmp_path/"cancel"), sys.executable)
    task = asyncio.create_task(box.execute_async("run", "import time\nopen('started.txt','w').close()\ntime.sleep(30)"))
    for _ in range(300):
        if Path(box.sandbox_path, "started.txt").exists():
            break
        await asyncio.sleep(0.01)
    else:
        pytest.fail("worker 未进入待取消业务")
    task.cancel()
    with pytest.raises(asyncio.CancelledError):
        await task
    assert processes and all(p.poll() is not None for p in processes)
    empty = CodeSandbox(str(tmp_path/"empty"), sys.executable, allowed_modules="")
    assert empty.run("print(2+3)")["returncode"] == 0
    assert empty.run("import math")["error_code"] == "module_not_allowed"
    invalid = CodeSandbox(str(tmp_path/"invalid"), sys.executable, allowed_modules="socket")
    assert invalid.execute("save", "never", "marker")["error_code"] == "configuration_error"
    secret = CodeSandbox(str(tmp_path/"secret"), sys.executable, allowed_env=["OPENAI_API_KEY"])
    assert secret.run("print('no')")["error_code"] == "configuration_error"
    limited = CodeSandbox(str(tmp_path/"limited"), sys.executable, allowed_modules="math,time")
    assert "math, time" in CodeSandboxTool(limited).description

    import venv
    environment = tmp_path / "without-dependencies"
    venv.EnvBuilder(with_pip=False).create(environment)
    interpreter = environment / ("Scripts/python.exe" if os.name == "nt" else "bin/python")
    missing = CodeSandbox(str(tmp_path/"missing"), str(interpreter))
    for operation in ("run", "save", "read", "list", "delete", "delete_dir", "run_file"):
        result = missing.execute(operation, "print('blocked')", "marker.txt")
        assert result["error_code"] == "sandbox_unavailable", result
    assert not Path(missing.sandbox_path, "marker.txt").exists()


def test_sandbox_protocol_limits_and_links(tmp_path: Path, monkeypatch) -> None:
    """真实子进程的损坏输出/超限结果报错, 文件链接不扩大业务范围"""
    from satrap.core.utils.sandbox import runner

    root=tmp_path/"sandbox"
    root.mkdir()
    sandbox=CodeSandbox(str(root),sys.executable)
    outside=tmp_path/"external"
    outside.mkdir()
    (outside/"private.txt").write_text("private",encoding="utf-8")
    try:
        (root/"linked").symlink_to(outside,target_is_directory=True)
    except OSError:
        pass
    else:
        assert sandbox.execute("read",path="linked/private.txt")["error_code"]=="path_outside_sandbox"
        assert sandbox.execute("save","bad","linked/private.txt")["error_code"]=="path_outside_sandbox"
        assert sandbox.execute("list")["error_code"]=="path_outside_sandbox"
        assert (outside/"private.txt").read_text(encoding="utf-8")=="private"

    fake=tmp_path/"worker.py"
    monkeypatch.setattr(runner,"__file__",str(tmp_path/"runner.py"))
    fake.write_text("print('{}')",encoding="utf-8")
    assert sandbox.run("print('no')")["error_code"]=="worker_failed"
    fake.write_text("print('x'*(8*1024*1024+1))",encoding="utf-8")
    assert sandbox.run("print('no')")["error_code"]=="output_limit_exceeded"


def test_sandbox_real_network_and_process_guards(tmp_path: Path) -> None:
    """可信子进程直接探测补丁, 不用模型导入拒绝代替真实禁网验证"""
    import json
    import socket
    import subprocess

    tcp = socket.socket()
    udp = socket.socket(type=socket.SOCK_DGRAM)
    tcp.bind(("127.0.0.1", 0))
    tcp.listen()
    tcp.settimeout(0.15)
    udp.bind(("127.0.0.1", 0))
    udp.settimeout(0.15)
    port, udp_port = tcp.getsockname()[1], udp.getsockname()[1]
    with socket.create_connection(("127.0.0.1", port)):
        accepted, _ = tcp.accept()
        accepted.close()
    udp.sendto(b"baseline", udp.getsockname())
    assert udp.recvfrom(1024)[0] == b"baseline"

    guard_dir = Path(__file__).parents[2] / "satrap/core/utils/sandbox"
    source = f"""
import sys,json,os
sys.path.insert(0,{str(guard_dir)!r})
from models import Violations, PolicyViolation
from guards import install
v=Violations()
installed=install(v)
import socket,ssl,subprocess,urllib.request,multiprocessing
assert issubclass(ssl.SSLSocket,socket.socket)
actions=[
    ("tcp",lambda:socket.create_connection(("127.0.0.1",{port}))),
    ("udp",lambda:socket.socket(type=socket.SOCK_DGRAM).sendto(b"blocked",("127.0.0.1",{udp_port}))),
    ("dns",lambda:socket.getaddrinfo("localhost",80)),
    ("http",lambda:urllib.request.urlopen("http://127.0.0.1:{port}/",timeout=1)),
    ("popen",lambda:subprocess.Popen(["satrap-does-not-exist"])),
    ("system",lambda:os.system("satrap-does-not-exist")),
    ("multiprocessing",lambda:multiprocessing.get_context("spawn").Process(target=print,args=("unexpected",)).start()),
]
if os.name=="nt":
    import _winapi
    actions.extend([
        ("startfile",lambda:os.startfile("satrap-does-not-exist.file")),
        ("CreateProcess",lambda:_winapi.CreateProcess(None,"satrap-does-not-exist",None,None,False,0,None,None,subprocess.STARTUPINFO())),
    ])
blocked=[]
for name,call in actions:
    try:
        call()
    except PolicyViolation as error:
        blocked.append([name,error.code])
installed.verify()
print(json.dumps(blocked))
"""
    try:
        r=subprocess.run([sys.executable,"-I","-S","-B","-X","utf8","-c",source],capture_output=True,timeout=15,text=True,encoding="utf-8")
        assert r.returncode==0, r.stderr
        blocked=dict(json.loads(r.stdout))
        assert set(blocked) >= {"tcp","udp","dns","http","popen","system","multiprocessing"}
        assert all(blocked[name]=="network_disabled" for name in ("tcp","udp","dns","http"))
        if os.name=="nt":
            assert blocked["startfile"]==blocked["CreateProcess"]=="process_not_allowed"
        with pytest.raises(socket.timeout):
            tcp.accept()
        with pytest.raises(socket.timeout):
            udp.recvfrom(1024)
        with socket.create_connection(("127.0.0.1",port)):
            accepted,_=tcp.accept()
            accepted.close()
        udp.sendto(b"host-still-online",udp.getsockname())
        assert udp.recvfrom(1024)[0]==b"host-still-online"
    finally:
        tcp.close()
        udp.close()


@pytest.fixture
def shell_env(tmp_path, monkeypatch):
    root = tmp_path / "workspace"
    root.mkdir()
    monkeypatch.setattr(coding, "WORKSPACE_ROOT", root)
    monkeypatch.setattr(coding, "_SYSTEM_PROTECTED_ROOT", root / ".satrap")
    engine = PermissionEngine(rules_file=tmp_path / "rules.json", log_file=tmp_path / "approval.jsonl")
    engine.set_plan_mode(True)
    session = SimpleNamespace(user_input_provider=lambda *args: "deny")
    bind_environment(session, replace(environment_snapshot(session), workspace_root=Path(root)))
    return root, engine, session

async def execute_shell(shell_env, asynchronous, command, shell="auto", cwd="", timeout=10):
    _, engine, session = shell_env
    tool = coding.AsyncShellTool(engine) if asynchronous else coding.ShellTool(engine)
    tool._bind(cast(Any, session))
    if isinstance(tool, coding.AsyncShellTool):
        return await tool.execute(command, timeout=timeout, shell=shell, cwd=cwd)
    return tool.execute(command, timeout=timeout, shell=shell, cwd=cwd)

@pytest.mark.parametrize("asynchronous", [False, True], ids=["sync", "async"])
async def test_shell_harness_read_control(shell_env, asynchronous):
    """同步与异步 shell 均只能读取工作区内文件"""
    root, engine, session = shell_env
    (root / "control.txt").write_text("audit-control", encoding="utf-8")
    reader = coding.AsyncReadFileTool() if asynchronous else coding.ReadFileTool()
    reader._bind(cast(Any, session))
    output = await reader.execute("control.txt") if isinstance(reader, coding.AsyncReadFileTool) else reader.execute("control.txt")
    assert "audit-control" in output
    engine.set_plan_mode(False)
    session.approval_provider = lambda request: ApprovalResult.for_request(request, "allow", "用户批准", "user")
    for shell in (("powershell", "cmd") if os.name == "nt" else ("bash", "sh")):
        output = await execute_shell(shell_env, asynchronous, "echo 审计中文-control", shell=shell)
        assert "审计中文-control" in output
        if shell == "cmd":
            command = f'"{sys.executable}" -X utf8 -c "print(\'中文 native\')"'
            assert "中文 native" in await execute_shell(shell_env, asynchronous, command, shell=shell)
            for symbol in ('&', '|', '>', '<'):
                quoted = f'"审计{symbol}quoted-output.txt"'
                output = await execute_shell(shell_env, asynchronous, f'echo {quoted}', shell=shell)
                assert quoted in output and "退出码" not in output
            assert not (root / "quoted-output.txt").exists()
            assert "退出码 7" in await execute_shell(shell_env, asynchronous, "exit /b 7", shell=shell)



def _native_python_command(source: str) -> str:
    """将真实 Python 场景编码为当前系统 Shell 可执行的固定命令"""
    import shlex
    payload = source.encode("utf-8").hex()
    args = [sys.executable, "-X", "utf8", "-c", 'import sys;exec(bytes.fromhex(sys.argv[1]))', payload]
    if os.name == "nt":
        return "& " + " ".join("'" + value.replace("'", "''") + "'" for value in args)
    return shlex.join(args)


@pytest.mark.parametrize("asynchronous", [False, True], ids=["sync", "async"])
async def test_shell_native_io_exit_and_output_budget(shell_env, monkeypatch, asynchronous):
    """真实 Shell 保留中文和工作目录, 关闭 stdin, 剥敏并限制输出内存"""
    root, engine, session = shell_env
    engine.set_plan_mode(False)
    engine.set_mode("full-access", persist=False)
    monkeypatch.setenv("SATRAP_NATIVE_TEST_TOKEN", "fake-native-secret")
    folder = root / "中文 空格"
    folder.mkdir()
    startup = root / "startup.sh"
    startup.write_text("touch unexpected-startup\n", encoding="utf-8")
    monkeypatch.setenv("BASH_ENV", str(startup))
    monkeypatch.setenv("ENV", str(startup))
    source = "import os,sys;from pathlib import Path;print('你好-shell');print(Path.cwd().name);print(repr(sys.stdin.read()));print(os.environ.get('SATRAP_NATIVE_TEST_TOKEN','removed'))"
    output = await execute_shell(shell_env, asynchronous, _native_python_command(source), cwd=folder.name)
    assert "你好-shell" in output and folder.name in output and "''" in output
    assert "removed" in output and "fake-native-secret" not in output
    assert not (folder / "unexpected-startup").exists()
    failed = await execute_shell(shell_env, asynchronous, _native_python_command("import sys;print('failure-中文',file=sys.stderr);sys.exit(7)"))
    assert "退出码 7" in failed and "failure-中文" in failed
    large = await execute_shell(shell_env, asynchronous, _native_python_command("import sys;sys.stdout.write('x'*10000000);sys.stderr.write('e'*10000000);print('tail-marker')"))
    assert "输出已截断" in large and large.endswith("tail-marker\n")
    assert len(large) < 20100
    timed = await execute_shell(shell_env, asynchronous, _native_python_command("import time;time.sleep(30)"), timeout=1)
    assert "执行超时" in timed


@pytest.mark.skipif(os.name == "nt", reason="POSIX 进程组回收场景")
@pytest.mark.parametrize("operation", ["timeout", "cancel", "parent_exit"])
async def test_shell_posix_descendant_cleanup(shell_env, operation):
    """真实异步工具超时及取消后等待进程组退出, 后代不再修改文件"""
    import asyncio
    import time
    root, engine, session = shell_env
    engine.set_plan_mode(False)
    engine.set_mode("full-access", persist=False)
    tool = coding.AsyncShellTool(engine)
    tool._bind(cast(Any, session))
    child = "import os,time;from pathlib import Path;Path('child.pid').write_text(str(os.getpid()));time.sleep(2);Path('escaped.txt').write_text('bad')"
    source = f"import subprocess,sys,time;subprocess.Popen([sys.executable,'-c',{child!r}]);time.sleep({0 if operation == 'parent_exit' else 30})"
    task = asyncio.create_task(tool.execute(_native_python_command(source), timeout=20 if operation == "cancel" else 1))
    deadline = time.monotonic() + 5
    while not (root / "child.pid").exists() and time.monotonic() < deadline and not task.done():
        await asyncio.sleep(0.02)
    assert (root / "child.pid").exists()
    if operation == "cancel":
        task.cancel()
        with pytest.raises(asyncio.CancelledError):
            await task
    else:
        assert "执行超时" in await task
    pid = int((root / "child.pid").read_text())
    proc = Path(f"/proc/{pid}/stat")
    deadline = time.monotonic() + 1
    if sys.platform.startswith("linux"):
        state = None
        while time.monotonic() < deadline:
            try:
                state = proc.read_text().rsplit(")", 1)[1].split()[0]
            except FileNotFoundError:
                state = None
                break
            if state == "Z":
                break
            await asyncio.sleep(0.02)
        assert state in {None, "Z"}
    await asyncio.sleep(2.1)
    assert not (root / "escaped.txt").exists()

@pytest.mark.parametrize("command", [
    "echo (New-Item audit-marker)",
    "echo audit\nNew-Item audit-marker",
    "git diff --output=audit-marker --no-index before.txt after.txt",
], ids=["parentheses", "newline", "git-output"])
async def test_shell_must_not_write_in_plan_mode(shell_env, command):
    root, _, _ = shell_env
    if command.startswith("git") and shutil.which("git") is None:
        pytest.skip("git-output 复现需要 Git")
    (root / "before.txt").write_text("before\n", encoding="utf-8")
    (root / "after.txt").write_text("after\n", encoding="utf-8")
    for asynchronous in (False, True):
        output = await execute_shell(shell_env, asynchronous, command)
        assert not (root / "audit-marker").exists(), f"未经批准创建了标记文件; output={output!r}"

@pytest.mark.parametrize("asynchronous", [False, True], ids=["sync", "async"])
async def test_shell_must_not_read_protected_token(shell_env, asynchronous):
    root, _, session = shell_env
    reader = coding.ReadFileTool()
    reader._bind(cast(Any, session))
    token = "audit-fake-token-does-not-authorize-anything"
    for token_path in (".satrap/api-token", ".satrap/credentials/api-token"):
        (root / token_path).parent.mkdir(parents=True, exist_ok=True)
        (root / token_path).write_text(token, encoding="utf-8")
        if "拒绝读取" not in reader.execute(token_path):
            raise RuntimeError("复现前置条件失败: 文件工具未拒绝受保护路径")
        output = await execute_shell(shell_env, asynchronous, ("Get-Content " if os.name == "nt" else "cat ") + token_path)
        assert token not in output, "Shell 输出了文件工具拒绝读取的伪造令牌"

@pytest.mark.parametrize("asynchronous", [False, True])
@pytest.mark.parametrize("answer", [None, "n", "all", "y"])
async def test_shell_requires_each_explicit_approval(tmp_path, monkeypatch, asynchronous, answer):
    engine = PermissionEngine(rules_file=tmp_path / "rules.json", log_file=tmp_path / "log.jsonl")
    calls = []
    approvals = []

    def provider(request):
        approvals.append(request)
        return ApprovalResult.for_request(request, "allow" if answer == "y" else "deny", "人工决定", "user")

    session = SimpleNamespace(approval_provider=provider if answer else None)

    bind_environment(session, replace(environment_snapshot(session), workspace_root=Path(tmp_path)))
    monkeypatch.setattr(coding, "_run_shell", lambda *args, **kwargs: calls.append(args) or "executed")
    tool = coding.AsyncShellTool(engine) if asynchronous else coding.ShellTool(engine)
    tool._bind(cast(Any, session))
    for _ in range(2):
        result = await tool.execute("echo harmless") if isinstance(tool, coding.AsyncShellTool) else tool.execute("echo harmless")
        assert (result == "executed") == (answer == "y")
    assert len(calls) == (2 if answer == "y" else 0)
    assert len(approvals) == (2 if answer else 0)
    assert len({request.request_id for request in approvals}) == len(approvals)
    assert all(json.loads(request.action_json)["cwd"] == str(tmp_path) for request in approvals)

@pytest.mark.parametrize("asynchronous", [False, True])
@pytest.mark.parametrize("change", ["plan", "workspace"])
async def test_shell_rechecks_scope_after_approval(tmp_path, monkeypatch, asynchronous, change):
    engine = PermissionEngine(rules_file=tmp_path / "rules.json", log_file=tmp_path / "log.jsonl")
    session = SimpleNamespace()
    bind_environment(session, replace(environment_snapshot(session), workspace_root=Path(tmp_path)))

    def provider(request):
        if change == "plan":
            engine.set_plan_mode(True)
        else:
            bind_environment(session, replace(environment_snapshot(session), workspace_root=Path(tmp_path.parent)))
        return ApprovalResult.for_request(request, "allow", "用户批准", "user")

    session.approval_provider = provider
    monkeypatch.setattr(coding, "_run_shell", lambda *args, **kwargs: pytest.fail("已失效授权不能执行"))
    tool = coding.AsyncShellTool(engine) if asynchronous else coding.ShellTool(engine)
    tool._bind(cast(Any, session))
    result = await tool.execute("echo audit") if isinstance(tool, coding.AsyncShellTool) else tool.execute("echo audit")
    assert "执行已取消" in result

def test_allowed_env_vars_parse_per_session():
    """allowlist 按会话配置解析, 空配置得到空集合, 不存在跨会话共享的模块级状态"""
    assert coding.parse_allowed_env_vars({"allowed_env_vars": "TEST_API_KEY,OTHER_TOKEN"}) == frozenset(
        {"TEST_API_KEY", "OTHER_TOKEN"}
    )
    assert coding.parse_allowed_env_vars({"allowed_env_vars": ""}) == frozenset()
    assert coding.parse_allowed_env_vars({}) == frozenset()
    assert not hasattr(coding, "ALLOWED_ENV_VARS")

@pytest.mark.parametrize("asynchronous", [False, True])
async def test_factory_environment_isolated_across_sessions(tmp_path, monkeypatch, asynchronous):
    import asyncio
    from satrap.edictum import SimpleSession, AsyncSimpleSession
    from satrap.core.utils import proc_env

    base = cast(type, AsyncSimpleSession if asynchronous else SimpleSession)   # 参数化基类

    class FakeSession(base):
        def __init__(self):
            self._wf = SimpleNamespace(llm=None, tools_manager=None)
            bind_environment(self, replace(environment_snapshot(self), workspace_root=Path(tmp_path)))
            self.approval_provider = lambda request: ApprovalResult.for_request(request, "allow", "用户批准", "user")

    def state(session, *, config=None):
        return {"engine": PermissionEngine(rules_file=tmp_path / "rules.json", log_file=tmp_path / "log.jsonl"), "todos": {}}

    monkeypatch.setattr(coding, "get_plugin_state", state)
    monkeypatch.setattr(proc_env.os, "environ", {"FIRST_TOKEN": "first", "LAST_API_KEY": "last"})
    def run(args, workdir, timeout, allow=frozenset(), **kwargs):
        return ",".join(sorted(proc_env.sanitized_child_env(allow=allow)))
    monkeypatch.setattr(coding, "_run_shell", run)
    configurations = [{"allowed_env_vars": "FIRST_TOKEN"}, {"allowed_env_vars": ""}, {"allowed_env_vars": "LAST_API_KEY"}]
    tools = [next(tool for tool in coding.get_tools(FakeSession(), config)
                  if isinstance(tool, (coding.ShellTool, coding.AsyncShellTool))) for config in configurations]
    async def execute(tool):
        if isinstance(tool, coding.AsyncShellTool):
            return await tool.execute("echo audit")
        return tool.execute("echo audit")
    assert [await execute(tool) for tool in [*tools, tools[0]]] == ["FIRST_TOKEN", "", "LAST_API_KEY", "FIRST_TOKEN"]
    if asynchronous:
        assert await asyncio.gather(*(execute(tool) for tool in tools)) == ["FIRST_TOKEN", "", "LAST_API_KEY"]

@pytest.mark.parametrize("asynchronous", [False, True], ids=["sync", "async"])
@pytest.mark.parametrize("tool_name", ["read_file", "write_file", "glob_files"])
async def test_factory_protected_directories_isolated(tmp_path, monkeypatch, asynchronous, tool_name):
    """真实工厂的保护配置在交错和并发调用中保持独立, 拒绝时不修改文件"""
    import asyncio
    from satrap.edictum import SimpleSession, AsyncSimpleSession

    base = cast(type, AsyncSimpleSession if asynchronous else SimpleSession)   # 参数化基类

    class FakeSession(base):
        def __init__(self):
            self._wf = SimpleNamespace(llm=None, tools_manager=None)
            bind_environment(self, replace(environment_snapshot(self), workspace_root=Path(tmp_path)))
            bind_environment(self, replace(environment_snapshot(self), sandbox_root=Path(tmp_path)))

    def state(session, *, config=None):
        return {"engine": PermissionEngine(rules_file=tmp_path / "rules.json",
                                          log_file=tmp_path / "log.jsonl"), "todos": {}}

    monkeypatch.setattr(coding, "get_plugin_state", state)
    configurations = [{"protected_dirs": " FIRST_PRIVATE, "},
                      {"protected_dirs": "second_private"}, {"protected_dirs": ""}, None]
    tools = [{tool.tool_name: tool for tool in coding.get_tools(FakeSession(), config)}
             for config in configurations]
    directories = ["first_private", "second_private", ".satrap", ".git", "node_modules"]
    for directory in directories:
        folder = tmp_path / directory
        folder.mkdir()
        (folder / "sample.txt").write_text("original-marker", encoding="utf-8")

    async def check(index):
        tool = tools[index][tool_name]
        for directory in directories:
            target = tmp_path / directory / f"case-{index}.txt"
            target.write_text("original-marker", encoding="utf-8")
            denied = directory in directories[2:] or (index < 2 and directory == directories[index])
            relative = f"{directory}/{target.name}"
            arguments = {
                "read_file": {"path": relative},
                "write_file": {"path": relative, "content": "changed-marker"},
                "edit_file": {"path": relative, "old": "original", "new": "changed"},
                "search_replace": {"path": relative, "replacements": [{"old": "original", "new": "changed"}]},
                "glob_files": {"pattern": f"{directory}/*.txt"},
                "grep_files": {"pattern": "marker", "path": directory},
            }[tool_name]
            result = await tool.execute(**arguments) if asynchronous else tool.execute(**arguments)
            if tool_name in {"glob_files", "grep_files"}:
                assert ("sample.txt" not in result) == denied, (index, directory, result)
            elif tool_name == "read_file":
                assert ("拒绝读取" in result) == denied, (index, directory, result)
                assert ("original-marker" not in result) == denied
            else:
                assert ("拒绝" in result) == denied, (index, directory, result)
                assert target.read_text(encoding="utf-8") == ("original-marker" if denied else "changed-marker")

    for index in [0, 1, 2, 3, 0]:
        await check(index)
    if asynchronous:
        await asyncio.gather(*(check(index) for index in range(4)))

def test_strips_suffix_matched_sensitive_vars(monkeypatch):
    """名称以 KEY/TOKEN/SECRET/PASSWORD 等为完整尾部分段的变量应被剥离"""
    monkeypatch.setenv("OPENAI_API_KEY", "sk-x")
    monkeypatch.setenv("SATRAP_API_TOKEN", "t")
    monkeypatch.setenv("AWS_SECRET_ACCESS_KEY", "s")
    monkeypatch.setenv("DB_PASSWORD", "p")
    monkeypatch.setenv("TOKEN", "bare")
    monkeypatch.setenv("service_api_key", "lowercase")

    env = sanitized_child_env()

    assert "OPENAI_API_KEY" not in env
    assert "SATRAP_API_TOKEN" not in env
    assert "AWS_SECRET_ACCESS_KEY" not in env
    assert "DB_PASSWORD" not in env
    assert "TOKEN" not in env
    assert "service_api_key" not in env

def test_keeps_normal_vars_containing_sensitive_prefix(monkeypatch):
    """含敏感词前缀但尾部分段不同的正常变量不应被误删"""
    monkeypatch.setenv("TOKENIZERS_PARALLELISM", "false")
    monkeypatch.setenv("SECRET_PROJECT_FOLDER", "x")
    monkeypatch.setenv("PATH", "C:/bin")

    env = sanitized_child_env()

    assert env["TOKENIZERS_PARALLELISM"] == "false"
    assert env["SECRET_PROJECT_FOLDER"] == "x"
    assert "PATH" in env

def test_allow_list_restores_explicit_vars(monkeypatch):
    """allow 显式放行的变量即使命中敏感特征也保留"""
    monkeypatch.setenv("TEST_SERVICE_TOKEN", "v")

    env = sanitized_child_env(allow=["TEST_SERVICE_TOKEN"])

    assert env["TEST_SERVICE_TOKEN"] == "v"

def test_overrides_win_over_stripped_values(monkeypatch):
    """overrides 追加的变量优先于剥敏结果"""
    monkeypatch.setenv("MY_TOKEN", "secret")

    env = sanitized_child_env(MY_TOKEN="safe")

    assert env["MY_TOKEN"] == "safe"

def test_allowlist_respects_platform_case_rules(monkeypatch):
    from satrap.core.utils import proc_env
    monkeypatch.setattr(proc_env.os, "environ", {"SERVICE_API_KEY": "test", "PATH": "bin"})
    monkeypatch.setattr(proc_env.os, "name", "nt")
    assert sanitized_child_env(allow=["service_api_key"])["SERVICE_API_KEY"] == "test"
    monkeypatch.setattr(proc_env.os, "name", "posix")
    assert "SERVICE_API_KEY" not in sanitized_child_env(allow=["service_api_key"])
    assert sanitized_child_env(allow=["SERVICE_API_KEY"])["SERVICE_API_KEY"] == "test"


@pytest.mark.parametrize("asynchronous", [False, True])
async def test_native_write_rechecks_live_platform_route(tmp_path, monkeypatch, asynchronous):
    """平台路由在人工等待中切换时, 获准的宿主写入仍取消"""
    from satrap.core.call_context import CallOrigin, bind_call_origin
    from satrap.core.platform import current_adapter_manager, set_current_adapter_manager
    from satrap.edictum import AsyncSimpleSession, SimpleSession
    from satrap.expend.plugins.satrap_coding.tools.paths import _execution_guard
    origin = CallOrigin("qq", "bot", "GroupMessage", "room", "actor", "message", "turn",
                        conversation_kind="group", agent_route_generation=1, group_route_generation=2)
    route = [1]
    adapter = SimpleNamespace(config=SimpleNamespace(enable=True), client_self_id="bot",
                              agent_route_store=SimpleNamespace(current_revision=lambda *args: route[0]),
                              group_route=lambda group: ({}, 2))
    manager = SimpleNamespace(get_adapter=lambda identity: adapter)
    previous = current_adapter_manager()
    set_current_adapter_manager(cast(Any, manager))
    engine = PermissionEngine(rules_file=tmp_path / "rules.json", log_file=tmp_path / "log.jsonl")
    session = SimpleNamespace()
    bind_environment(session, replace(environment_snapshot(session), workspace_root=Path(tmp_path)))
    def approve(request):
        route[0] = 2
        return ApprovalResult.for_request(request, "allow", "用户批准", "user")
    session.approval_provider = approve
    tool = coding.AsyncWriteFileTool(engine) if asynchronous else coding.WriteFileTool(engine)
    tool._bind(cast(Any, session))
    try:
        with bind_call_origin(origin):
            result = tool.execute("result.txt", "must-not-write")
            result = await result if asynchronous else result
        assert "来源 Agent 路由已经变化" in result
        assert not (tmp_path / "result.txt").exists()
    finally:
        set_current_adapter_manager(previous)
